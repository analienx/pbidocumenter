"""GitHub-runner-only source fidelity and rendered Word-page acceptance evidence.

LibreOffice is a repeatable PDF rendering oracle, NOT an assertion of Word-exact
pagination. Page images and JSON metrics are uploaded for an independent review.
"""

from __future__ import annotations

import argparse
import json
import subprocess
import tempfile
from pathlib import Path

import fitz
from copiloter import PbipProject, build_project_summary
from docx import Document

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE = PROJECT_ROOT / "examples" / "contoso-retail"


def _normalize(value: str) -> str:
    return " ".join(value.split())


def _inventory_table(doc: Document, first_cells: list[str]) -> list[list[str]]:
    for table in doc.tables:
        if not table.rows or len(table.rows[0].cells) < len(first_cells):
            continue
        header = [_normalize(c.text) for c in table.rows[0].cells[: len(first_cells)]]
        if header == first_cells:
            return [[_normalize(c.text) for c in row.cells] for row in table.rows[1:]]
    raise AssertionError(f"Missing document inventory table with headers {first_cells!r}")


def _verify_source_fidelity(docx: Path, source: dict) -> dict:
    doc = Document(docx)
    sm = source["semantic_model"]
    rpt = source["report"]

    table_rows = _inventory_table(doc, ["Table", "Role", "Storage Mode"])
    measure_rows = _inventory_table(doc, ["#", "Measure Name", "Source Table"])
    rel_rows = _inventory_table(doc, ["#", "From Table", "From Column", "To Table", "To Column"])
    page_rows = _inventory_table(doc, ["Page", "Page Type", "Total Visuals"])

    expected_tables = sorted(str(t["name"]) for t in sm["tables"])
    actual_tables = sorted(row[0] for row in table_rows)
    assert actual_tables == expected_tables, f"Table inventory mismatch: {actual_tables} != {expected_tables}"

    expected_measures = sorted(
        (str(t["name"]), str(m["name"])) for t in sm["tables"] for m in t.get("measures", [])
    )
    actual_measures = sorted((row[2], row[1]) for row in measure_rows)
    assert actual_measures == expected_measures, "Measure names/table binding do not match source"

    expected_relationships = sorted(
        (
            str(rel.get("from_table", "")),
            str(rel.get("from_column", "")),
            str(rel.get("to_table", "")),
            str(rel.get("to_column", "")),
        )
        for rel in sm["relationships"]
    )
    actual_relationships = sorted((row[1], row[2], row[3], row[4]) for row in rel_rows)
    assert actual_relationships == expected_relationships, "Relationship endpoint inventory differs from source"

    expected_pages = sorted(str(pg["display_name"]) for pg in rpt["pages"])
    actual_pages = sorted(row[0] for row in page_rows)
    assert actual_pages == expected_pages, "Report pages differ from PBIP source"
    assert sum(int(row[2]) for row in page_rows) == rpt["total_visuals"], "Page visual totals are incorrect"

    content = _normalize("\n".join(
        [par.text for par in doc.paragraphs]
        + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    ))
    assert "No budget/forecast comparisons" not in content, "Invented budget exclusion survives"

    assert len(doc.inline_shapes) >= 1, "Missing Word diagrams/wireframes"
    return {
        "tables_verified": len(actual_tables),
        "measures_verified": len(actual_measures),
        "relationships_verified": len(actual_relationships),
        "pages_verified": len(actual_pages),
        "visuals_verified": rpt["total_visuals"],
        "inline_shapes": len(doc.inline_shapes),
        "structured_headings": sum(p.style.name.startswith("Heading") for p in doc.paragraphs),
    }


def _render_and_check(docx: Path, source: dict, destination: Path) -> dict:
    pdf_path = destination / f"{docx.stem}.pdf"
    with tempfile.TemporaryDirectory(prefix="pbipdoc-lo-profile-") as home:
        profile = Path(home).as_uri()
        result = subprocess.run(
            [
                "libreoffice",
                f"-env:UserInstallation={profile}",
                "--headless",
                "--convert-to",
                "pdf:writer_pdf_Export",
                "--outdir",
                str(destination),
                str(docx),
            ],
            capture_output=True,
            text=True,
            timeout=150,
            check=False,
        )
    assert result.returncode == 0 and pdf_path.is_file(), (
        f"LibreOffice failed to render {docx.name}: {result.stdout[-1000:]} {result.stderr[-1000:]}"
    )

    page_stats = []
    pdf_text_parts = []
    with fitz.open(pdf_path) as pdf:
        assert 3 <= len(pdf) <= 200, f"Suspicious number of PDF pages: {len(pdf)}"
        for number, page in enumerate(pdf, start=1):
            text = page.get_text(sort=True)
            pdf_text_parts.append(text)
            blocks = page.get_text("blocks")
            escaped = []
            for block in blocks:
                x0, y0, x1, y1 = block[:4]
                if x0 < -3 or y0 < -3 or x1 > page.rect.width + 3 or y1 > page.rect.height + 3:
                    escaped.append([round(x0, 1), round(y0, 1), round(x1, 1), round(y1, 1)])
            # Pure-white/empty pages are a rendering defect, except if the
            # entire page consists of intentional diagrams/drawings.
            empty = not text.strip() and not page.get_drawings() and not page.get_images()
            assert not empty, f"Unexpected completely blank PDF page: {docx.name} page {number}"
            preview = destination / f"{docx.stem}-page-{number:03}.jpg"
            pix = page.get_pixmap(matrix=fitz.Matrix(0.9, 0.9), alpha=False)
            pix.save(preview, jpg_quality=78)
            page_stats.append(
                {
                    "page": number,
                    "text_chars": len(text.strip()),
                    "drawings": len(page.get_drawings()),
                    "images": len(page.get_images()),
                    "out_of_page_blocks": escaped,
                    "preview": preview.name,
                }
            )

    rendered_text = _normalize("\n".join(pdf_text_parts))
    for table in source["semantic_model"]["tables"]:
        name = _normalize(str(table["name"]))
        assert name in rendered_text, f"Source table missing from rendered PDF: {name}"
    for page in source["report"]["pages"]:
        name = _normalize(str(page["display_name"]))
        assert name in rendered_text, f"Source page missing from rendered PDF: {name}"
    return {
        "render_engine": "LibreOffice headless (not Word-exact)",
        "pdf": pdf_path.name,
        "total_pages": len(page_stats),
        "total_rendered_text_chars": len(rendered_text),
        "low_text_pages": [p["page"] for p in page_stats if p["text_chars"] < 35],
        "outside_page_blocks": sum(len(p["out_of_page_blocks"]) for p in page_stats),
        "pages": page_stats,
    }


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--docx-dir", required=True, type=Path)
    ap.add_argument("--render-dir", required=True, type=Path)
    args = ap.parse_args()
    destination = args.render_dir.resolve()
    destination.mkdir(parents=True, exist_ok=True)

    project = PbipProject(SOURCE)
    assert project.is_valid, "Public Contoso source is not recognized as PBIP"
    source = build_project_summary(project)
    results = {"source_project": str(SOURCE.relative_to(PROJECT_ROOT)), "documents": {}}
    for mode in ("default", "full"):
        docx = args.docx_dir.resolve() / f"contoso-{mode}.docx"
        assert docx.is_file(), f"Missing installed-wheel output: {docx}"
        inventory = _verify_source_fidelity(docx, source)
        render = _render_and_check(docx, source, destination)
        results["documents"][mode] = {"inventory": inventory, "render": render}
        print(
            f"ACCEPTED {mode}: {inventory['pages_verified']} report pages, "
            f"{inventory['tables_verified']} tables, {inventory['measures_verified']} measures, "
            f"{render['total_pages']} rendered PDF pages, "
            f"{render['outside_page_blocks']} out-of-page blocks"
        )

    report_path = destination / "word-render-acceptance.json"
    report_path.write_text(json.dumps(results, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"Acceptance evidence: {report_path}")


if __name__ == "__main__":
    main()
