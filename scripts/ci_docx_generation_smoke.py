"""Installed-wheel end-to-end DOCX generation smoke for GitHub-hosted CI."""

from __future__ import annotations

import argparse
import subprocess
import sys
import tempfile
from pathlib import Path
from zipfile import ZipFile

from docx import Document

SOURCE_ROOT = Path(__file__).resolve().parents[1]
EXAMPLE = SOURCE_ROOT / "examples" / "contoso-retail"


def _assert_valid_docx(path: Path) -> None:
    assert path.is_file(), f"Missing generated DOCX: {path}"
    assert path.stat().st_size > 10_000, f"Suspiciously small DOCX: {path}"
    with ZipFile(path) as archive:
        assert archive.testzip() is None, f"Corrupt DOCX archive: {path}"
        assert "word/document.xml" in archive.namelist()

    doc = Document(path)
    from docx.oxml.ns import qn

    linked = {n.get(qn("w:anchor")) for n in doc.element.body.xpath(".//w:hyperlink")}
    targets = {n.get(qn("w:name")) for n in doc.element.body.xpath(".//w:bookmarkStart")}
    required = {
        "toc_control", "toc_scope", "toc_operations", "toc_technical", "toc_support", "toc_revision"
    }
    assert required <= linked & targets, "Generated DOCX has no working linked contents"
    assert len(doc.tables) >= 3, f"No substantive specification tables: {path}"
    assert any(p.style.name.startswith("Heading") for p in doc.paragraphs), f"No Word heading styles: {path}"
    content = "\n".join(
        [p.text for p in doc.paragraphs]
        + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    )
    for expected in ("Fact Sales", "Dim Customer", "Dim Date"):
        assert expected in content, f"Semantic object '{expected}' missing from {path}"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()

    destination = args.output_dir.resolve()
    destination.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="pbipdoc-installed-wheel-") as isolated:
        workdir = Path(isolated)
        for mode in ("default", "full"):
            output = destination / f"contoso-{mode}.docx"
            process = subprocess.run(
                [
                    sys.executable,
                    "-m",
                    "pbip_documenter",
                    str(EXAMPLE),
                    "--mode",
                    mode,
                    "--output",
                    str(output),
                ],
                cwd=workdir,
                capture_output=True,
                text=True,
                check=False,
            )
            assert process.returncode == 0, (
                f"Installed CLI {mode} generation failed ({process.returncode}):\n"
                f"{process.stdout[-2500:]}\n{process.stderr[-2500:]}"
            )
            _assert_valid_docx(output)
            print(f"OK: {output.name}, {output.stat().st_size:,} bytes")

        process = subprocess.run(
            [sys.executable, "-m", "pbip_documenter", str(EXAMPLE)],
            cwd=workdir,
            capture_output=True,
            text=True,
            check=False,
        )
        assert process.returncode == 0, f"Default-path CLI failed: {process.stderr[-1000:]}"
        defaults = sorted((workdir / "Exported Documents").rglob("*.docx"))
        assert len(defaults) == 1, f"Expected one user-visible default output, found {len(defaults)}"
        _assert_valid_docx(defaults[0])
        print(f"OK: default working-directory output: {defaults[0].relative_to(workdir)}")


if __name__ == "__main__":
    main()
