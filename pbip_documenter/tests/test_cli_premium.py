"""Premium CLI behavior and safety contracts, using only disposable model fixtures."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pytest
from docx import Document

from copiloter import PbipProject
from pbip_documenter import bootstrap, cli
from pbip_documenter.cli_validation import source_fingerprint


def model_project(root: Path, name: str) -> Path:
    project = root / name
    tables = project / f"{name}.SemanticModel" / "definition" / "tables"
    tables.mkdir(parents=True)
    (tables / "Fact.tmdl").write_text("table Fact\n\tcolumn ID\n\t\tdataType: int64\n", encoding="utf-8")
    return project


def command(capsys: pytest.CaptureFixture[str], argv: list[str]) -> tuple[int, dict[str, Any]]:
    code = bootstrap.main(argv + ["--json"])
    output = capsys.readouterr()
    assert not output.err, output.err
    assert output.out.count("\n") == 1, output.out[:300]
    result = json.loads(output.out)
    assert result["exit_code"] == code
    return code, result


def stub_doc(*_args: Any, **_kwargs: Any) -> Any:
    doc = Document()
    doc.add_heading("Native Word specification", level=1)
    return doc


def test_dry_run_is_read_only_with_explicit_destination(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], monkeypatch: pytest.MonkeyPatch
) -> None:
    project = model_project(tmp_path, "ReportOne")
    monkeypatch.setattr(cli, "build_doc", lambda *_a, **_k: pytest.fail("dry-run must not render"))
    destination = tmp_path / "not_created" / "spec.docx"
    code, data = command(capsys, [
        "generate", str(project), "--dry-run", "--manifest", "--output", str(destination)
    ])
    assert code == 0 and data["dry_run"] is True
    assert data["planned_outputs"] == [str(destination.resolve())]
    assert data["outputs"] == [] and data["artifacts"] == []
    assert len(data["projects"][0]["source_sha256"]) == 64
    assert not destination.parent.exists()


def test_multi_project_preview_and_generation_with_evidence(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], monkeypatch: pytest.MonkeyPatch
) -> None:
    root = tmp_path / "projects"
    model_project(root, "Alpha")
    model_project(root, "Beta")
    outdir = tmp_path / "results"
    monkeypatch.setattr(cli, "build_doc", stub_doc)

    code, preview = command(capsys, [
        "generate", str(root), "--output-dir", str(outdir), "--manifest", "--dry-run"
    ])
    assert code == 0 and len(preview["planned_outputs"]) == 2
    assert all(not Path(p).exists() for p in preview["planned_outputs"])
    assert not outdir.exists()

    code, data = command(capsys, [
        "generate", str(root), "--output-dir", str(outdir), "--manifest", "--no-overwrite"
    ])
    assert code == 0
    assert data["outputs"] == preview["planned_outputs"]
    assert len(data["artifacts"]) == 2
    for evidence in data["artifacts"]:
        doc = Path(evidence["document"])
        sidecar = Path(evidence["manifest"])
        assert doc.is_file() and sidecar.is_file()
        stored = json.loads(sidecar.read_text(encoding="utf-8"))
        assert evidence["source_sha256"] == stored["source_sha256"]
        assert evidence["document_sha256"] == stored["document_sha256"]
        assert evidence["bytes"] == doc.stat().st_size
        assert stored["generator"] == "pbip-documenter"
        assert stored["schema_version"] == "1.0"
        assert len(stored["source_sha256"]) == 64

    original = {p: Path(p).read_bytes() for p in data["outputs"]}
    code, failure = command(capsys, [
        "generate", str(root), "--output-dir", str(outdir), "--manifest", "--no-overwrite"
    ])
    assert code == 2
    assert all(d["code"] == "OUTPUT_EXISTS" for d in failure["diagnostics"])
    assert {p: Path(p).read_bytes() for p in data["outputs"]} == original


def test_conflicting_destination_flags_rejected(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    project = model_project(tmp_path, "One")
    code, result = command(capsys, [
        "generate", str(project), "--output", str(tmp_path / "a.docx"),
        "--output-dir", str(tmp_path / "folder"), "--dry-run"
    ])
    assert code == 2
    assert result["diagnostics"][0]["code"] == "OUTPUT_CONFLICT"


def test_hint_from_invalid_tmdl(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    project = model_project(tmp_path, "Broken")
    table = next(project.rglob("Fact.tmdl"))
    table.write_text("not a table definition", encoding="utf-8")
    code, result = command(capsys, ["validate", str(project), "--strict"])
    assert code == 3
    findings = result["projects"][0]["diagnostics"]
    assert any(x["code"] == "TMDL_TABLE_UNPARSED" and "hint" in x for x in findings)
    assert result["projects"][0]["diagnostic_counts"]["error"] >= 1


def test_fingerprint_scoped_to_source_files(
    tmp_path: Path
) -> None:
    project = model_project(tmp_path, "Sources")
    selected = PbipProject(project)
    first = source_fingerprint(selected)
    (project / "notes.json").write_text('{"comment":"not part of source"}', encoding="utf-8")
    (project / "evidence.json").write_text('{"other":"content"}', encoding="utf-8")
    assert first == source_fingerprint(selected), "Unrelated files must not change source revision"
    table = next(project.rglob("Fact.tmdl"))
    table.write_text("table Fact\n\tcolumn ID\n\tcolumn Name\n", encoding="utf-8")
    assert first != source_fingerprint(selected)


def test_symlink_cannot_fingerprint_external_file(tmp_path: Path) -> None:
    project = model_project(tmp_path, "Source")
    external = tmp_path / "external-secret.json"
    external.write_text('{"sensitive":"TOKEN"}', encoding="utf-8")
    digest = source_fingerprint(PbipProject(project))
    link = project / "Source.SemanticModel" / "definition" / "external.json"
    try:
        link.symlink_to(external)
    except (OSError, NotImplementedError):
        pytest.skip("File symlinks not supported by this Windows runner")
    assert source_fingerprint(PbipProject(project)) == digest


def test_no_overwrite_is_not_a_noop(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], monkeypatch: pytest.MonkeyPatch
) -> None:
    project = model_project(tmp_path, "Safe")
    out = tmp_path / "saved.docx"
    out.write_bytes(b"DO NOT REPLACE")
    monkeypatch.setattr(cli, "build_doc", lambda *_a, **_k: pytest.fail("must refuse before building"))
    code, result = command(capsys, [
        "generate", str(project), "--output", str(out), "--no-overwrite"
    ])
    assert code == 2
    assert "OUTPUT_EXISTS" in [x["code"] for x in result["diagnostics"]]
    assert out.read_bytes() == b"DO NOT REPLACE"

def test_invalid_explicit_template_rejected_early(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], monkeypatch: pytest.MonkeyPatch
) -> None:
    project = model_project(tmp_path, "Template")
    template = tmp_path / "fake.docx"
    template.write_text("not a Word ZIP", encoding="utf-8")
    monkeypatch.setattr(cli, "build_doc", lambda *_a, **_k: pytest.fail("must reject template before build"))
    code, result = command(capsys, [
        "generate", str(project), "--template", str(template), "--dry-run"
    ])
    assert code == 2
    assert result["diagnostics"][0]["code"] == "TEMPLATE_INVALID"


def test_invalid_logo_and_output_name_rejected(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    project = model_project(tmp_path, "Names")
    logo = tmp_path / "fake.pdf"
    logo.write_bytes(b"not a logo")
    code, result = command(capsys, [
        "generate", str(project), "--logo", str(logo), "--output", str(tmp_path / "bad.txt"), "--dry-run"
    ])
    assert code == 2
    assert {x["code"] for x in result["diagnostics"]} == {"LOGO_INVALID", "OUTPUT_INVALID"}


def test_atomic_no_clobber_even_when_second_process_creates_output(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], monkeypatch: pytest.MonkeyPatch
) -> None:
    project = model_project(tmp_path, "Race")
    destination = tmp_path / "race.docx"

    class SimulatedConcurrentDocument:
        def save(self, path: str) -> None:
            Path(path).write_bytes(b"NEW-DOCUMENT")
            destination.write_bytes(b"OTHER-PROCESS-DOCUMENT")

    monkeypatch.setattr(cli, "build_doc", lambda *_a, **_k: SimulatedConcurrentDocument())
    code, result = command(capsys, [
        "generate", str(project), "--output", str(destination), "--no-overwrite"
    ])
    assert code != 0
    assert result["status"] == "error"
    assert destination.read_bytes() == b"OTHER-PROCESS-DOCUMENT"
    assert not list(tmp_path.glob(".pbipdoc-*.docx"))

def test_input_mutation_during_word_render_cannot_publish_stale_document(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], monkeypatch: pytest.MonkeyPatch
) -> None:
    project = model_project(tmp_path, "Mutable")
    output = tmp_path / "already-valid.docx"
    output.write_bytes(b"PREVIOUS-VALID-WORD-FILE")
    table = next(project.rglob("Fact.tmdl"))

    class MutatingDocument:
        def save(self, path: str) -> None:
            Path(path).write_bytes(b"INCOMPLETE-DOCX")
            table.write_text("table Fact\n\tcolumn SomethingElse", encoding="utf-8")

    monkeypatch.setattr(cli, "build_doc", lambda *_a, **_k: MutatingDocument())
    code, result = command(capsys, ["generate", str(project), "--output", str(output)])
    assert code == 3
    assert any(item["code"] == "SOURCE_CHANGED" for item in result["diagnostics"])
    assert output.read_bytes() == b"PREVIOUS-VALID-WORD-FILE"
    assert not list(tmp_path.glob(".pbipdoc-*.docx"))
