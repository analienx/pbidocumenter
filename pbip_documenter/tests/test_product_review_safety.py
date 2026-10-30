"""Full-product code-review regressions for data fidelity, safe output and downloads."""

from __future__ import annotations

import json
from pathlib import Path
from unittest.mock import Mock

import pytest
from docx import Document

from copiloter import PbipProject, build_project_summary
from pbip_documenter import cli
from pbip_documenter.downloads.service import DownloadService


def make_pbir(tmp_path: Path, *, corrupt_visual: bool = False, corrupt_page: bool = False) -> Path:
    root = tmp_path / "Review.Report"
    visual = root / "definition" / "pages" / "TestPage" / "visuals" / "ChartA"
    visual.mkdir(parents=True)
    (root / ".platform").write_text(json.dumps({"metadata": {"displayName": "Review"}}), encoding="utf-8")
    (root / "definition.pbir").write_text(
        json.dumps({"datasetReference": {"byPath": {"path": "../Model"}}}),
        encoding="utf-8",
    )
    (visual.parent.parent / "page.json").write_text(
        "{invalid" if corrupt_page else json.dumps({"name": "TestPage", "displayName": "Test Page"}),
        encoding="utf-8",
    )
    (visual / "visual.json").write_text(
        "{invalid" if corrupt_visual else json.dumps({"name": "ChartA", "visual": {"visualType": "barChart"}}),
        encoding="utf-8",
    )
    return root


@pytest.mark.parametrize("corrupt_visual,corrupt_page", [(True, False), (False, True)])
def test_malformed_existing_pbir_fails_closed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, corrupt_visual: bool, corrupt_page: bool
) -> None:
    project = make_pbir(tmp_path, corrupt_visual=corrupt_visual, corrupt_page=corrupt_page)
    summary = build_project_summary(PbipProject(project))
    assert any(str(w).startswith("ERROR:") for w in summary["report"]["warnings"])
    build = Mock(side_effect=AssertionError("Must not build from incomplete report"))
    monkeypatch.setattr(cli, "build_doc", build)
    output = tmp_path / "bad.docx"
    assert cli.main([str(project), "--output", str(output)]) != 0
    assert not output.exists()
    build.assert_not_called()


def test_valid_pbir_visual_still_succeeds(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    project = make_pbir(tmp_path)
    summary = build_project_summary(PbipProject(project))
    assert summary["report"]["total_visuals"] == 1
    assert not any(str(w).startswith("ERROR:") for w in summary["report"]["warnings"])
    output = tmp_path / "valid.docx"
    monkeypatch.setattr(cli, "build_doc", lambda *_a, **_k: Document())
    assert cli.main([str(project), "--output", str(output)]) == 0
    assert output.is_file()


def test_template_cannot_be_overwritten_as_output(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    project = make_pbir(tmp_path)
    template = tmp_path / "template.docx"
    doc = Document()
    doc.add_paragraph("KEEP")
    doc.save(str(template))
    original = template.read_bytes()
    build = Mock(side_effect=AssertionError("Must reject unsafe output before build"))
    monkeypatch.setattr(cli, "build_doc", build)
    assert cli.main([str(project), "--template", str(template), "--output", str(template)]) == 2
    assert template.read_bytes() == original
    build.assert_not_called()


def test_browser_download_stabilization_respects_deadline(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    service = DownloadService.__new__(DownloadService)
    service.config = Mock(poll_seconds=0, timeout_seconds=1)
    partial = tmp_path / "download.crdownload"
    partial.write_bytes(b"unfinished")
    # Model the monotonic clock advancing between attempts, avoiding real sleeps.
    ticks = iter([0.0, 0.1, 0.2, 0.8, 1.1, 1.2])
    monkeypatch.setattr("pbip_documenter.downloads.service.time.monotonic", lambda: next(ticks))
    monkeypatch.setattr("pbip_documenter.downloads.service.time.sleep", lambda _seconds: None)
    with pytest.raises(TimeoutError, match="stabilize"):
        service._wait_until_file_stable(partial, stable_checks=2, deadline=1.0)


@pytest.mark.parametrize("name", ["../escaped.json", "../../escape.json", "/absolute.json"])
def test_cache_rejects_unconfined_destinations(tmp_path: Path, name: str) -> None:
    service = DownloadService.__new__(DownloadService)
    service.cache_dir = tmp_path / "cache"
    service.cache_dir.mkdir()
    source = tmp_path / "downloaded.json"
    source.write_text('{"test": true}', encoding="utf-8")
    with pytest.raises(ValueError, match="cache directory"):
        service._cache_file(source, name)
    assert not (tmp_path / "escaped.json").exists()


def test_cache_allows_valid_name(tmp_path: Path) -> None:
    service = DownloadService.__new__(DownloadService)
    service.cache_dir = tmp_path / "cache"
    service.cache_dir.mkdir()
    source = tmp_path / "downloaded.json"
    source.write_text('{"test": true}', encoding="utf-8")
    cached = service._cache_file(source, "inventory")
    assert cached == service.cache_dir / "inventory.json"
    assert cached.read_bytes() == source.read_bytes()

def test_failed_document_save_keeps_previous_good_output(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    project = make_pbir(tmp_path)
    output = tmp_path / "existing.docx"
    output.write_bytes(b"PREVIOUS VALID DOCUMENT")
    class FailedDoc:
        def save(self, path: str) -> None:
            Path(path).write_bytes(b"INCOMPLETE")
            raise OSError("simulated write failure")
    monkeypatch.setattr(cli, "build_doc", lambda *_args, **_kwargs: FailedDoc())
    with pytest.raises(OSError, match="simulated write failure"):
        cli.main([str(project), "--output", str(output)])
    assert output.read_bytes() == b"PREVIOUS VALID DOCUMENT"
    assert not list(tmp_path.glob(".pbipdoc-*.docx"))


def test_live_connection_string_secrets_not_exported(tmp_path: Path) -> None:
    project = make_pbir(tmp_path)
    (project / "definition.pbir").write_text(
        json.dumps({
            "datasetReference": {
                "byConnection": {
                    "connectionType": "pbiServiceXmlaStyleLive",
                    "connectionString": (
                        'Data Source=powerbi://example.test;Initial Catalog=SafeCatalog;'
                        'Password=NEVER_EXPORT_THIS_SECRET;Access Token=TEST_SECRET;'
                    ),
                },
            },
        }),
        encoding="utf-8",
    )
    summary = build_project_summary(PbipProject(project))
    serialized = json.dumps(summary, default=str)
    assert "NEVER_EXPORT_THIS_SECRET" not in serialized
    assert "TEST_SECRET" not in serialized
    details = summary["report"]["dataset_connection"]
    assert "connection_string" not in details
    assert details["catalog"] == "SafeCatalog"
    assert details["data_source"] == "powerbi://example.test"

def test_word_header_footer_references_are_document_local(tmp_path: Path) -> None:
    from docx.oxml.ns import qn

    from pbip_documenter.services.document_service import (
        _begin_schema_landscape_section,
        _end_schema_portrait_section,
    )

    def make_document(name: str, add_header: bool) -> tuple[Path, list[str]]:
        doc = Document()
        if add_header:
            doc.sections[0].header.paragraphs[0].text = "PRIVATE HEADER FOR " + name
        _begin_schema_landscape_section(doc)
        _end_schema_portrait_section(doc)
        refs = [
            ref
            for tag in ("headerReference", "footerReference")
            for node in doc.sections[-1]._sectPr.findall(qn(f"w:{tag}"))
            if (ref := node.get(qn("r:id"))) is not None
        ]
        destination = tmp_path / f"{name}.docx"
        doc.save(str(destination))
        return destination, refs

    doc_a, refs_a = make_document("A", True)
    doc_b, refs_b = make_document("B", False)
    assert refs_a
    assert not refs_b, "The second document must not inherit another document's header/footer IDs"
    assert "PRIVATE HEADER FOR A" not in doc_b.read_bytes().decode("utf-8", errors="ignore")
    assert doc_a.is_file() and doc_b.is_file()
