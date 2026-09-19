"""Independent visual evidence cannot be inferred from structural PBIR validity."""

from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document
from docx.shared import Pt

from pbip_documenter.visual_quality.document import inspect_document
from pbip_documenter.visual_quality.evidence import digest, review_template, verify_review
from pbip_documenter.visual_quality.policy import REQUIRED
from pbip_documenter.visual_quality.runner import audit, iterate, request, snapshot

ROOT = Path(__file__).parents[2] / "examples" / "contoso-retail"
REPORT = ROOT / "Contoso Retail.Report"
RENDERS = ROOT / "screenshots"


def _approved() -> tuple[dict, list[dict], str]:
    state = snapshot("report", REPORT, RENDERS)
    review = review_template("report", state["source_sha256"], state["pages"], "fixer")
    review["reviewer"]["id"] = "separate-visual-reviewer"
    for page in review["pages"]:
        for answer in page["observations"]:
            answer.update(status="pass", reason="This observation was explicitly checked on the source-bound rendered image.")
    return review, state["pages"], state["source_sha256"]


def test_policy_is_complete_and_requires_25_page_judgments() -> None:
    review, pages, sha = _approved()
    assert len(pages) == 5 and len(REQUIRED["report"]) == 25
    assert len(review["pages"][0]["observations"]) == 25
    assert not verify_review("report", sha, pages, review, "fixer")


def test_missing_review_or_same_executor_fails() -> None:
    assert any(f["rule"] == "independent_review_missing" for f in
               audit("report", REPORT, RENDERS, None, "fixer")["findings"])
    review, pages, sha = _approved()
    review["reviewer"]["id"] = "fixer"
    assert any(f["rule"] == "independent_reviewer_required" for f in
               verify_review("report", sha, pages, review, "fixer"))


def test_each_observation_needs_fresh_image_and_actionable_failure() -> None:
    review, pages, sha = _approved()
    review["pages"][0]["image_sha256"] = "stale"
    assert any(f["rule"] == "review_image_stale" for f in
               verify_review("report", sha, pages, review, "fixer"))
    review, pages, sha = _approved()
    finding = review["pages"][0]["observations"][0]
    finding.update(status="fail", severity="high", region=[0.1, 0.2, 0.6, 0.7],
                   proposed_fix="Increase the title and reduce clutter.",
                   reason="The header visually competes with the primary metric at normal page size.")
    assert any(f["rule"] == "visual_defect" and f["severity"] == "high"
               for f in verify_review("report", sha, pages, review, "fixer"))
    finding["region"] = None
    assert any(f["rule"] == "issue_lacks_location_or_repair" for f in
               verify_review("report", sha, pages, review, "fixer"))


def test_mandatory_criteria_cannot_be_marked_inapplicable() -> None:
    review, pages, sha = _approved()
    review["pages"][1]["observations"][0]["status"] = "not_applicable"
    assert any(f["rule"] == "mandatory_observation_skipped" for f in
               verify_review("report", sha, pages, review, "fixer"))


def test_word_ooxml_preflight_is_not_a_rendered_review(tmp_path: Path) -> None:
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("A meaningful paragraph of document content.").font.size = Pt(7)
    path = tmp_path / "output.docx"
    doc.save(path)
    state = inspect_document(path)
    assert state["surface"] == "document" and state["source_sha256"] == digest(path)
    assert not state["render_pass"]
    assert any(f["rule"] == "small_direct_font" for f in state["static_findings"])


def test_word_image_manifest_needs_exact_source_and_all_pages(tmp_path: Path) -> None:
    path = tmp_path / "output.docx"
    Document().save(path)
    renders = tmp_path / "rendered"
    renders.mkdir()
    example = RENDERS / "01-executive-performance.png"
    shutil.copyfile(example, renders / "page-0001.png")
    manifest = {"source_sha256": digest(path), "page_order": ["page-0001.png"],
                "files": {"page-0001.png": digest(renders / "page-0001.png")}}
    (renders / "capture-manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    assert inspect_document(path, renders)["render_pass"]
    manifest["source_sha256"] = "different"
    (renders / "capture-manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    assert not inspect_document(path, renders)["render_pass"]


def test_iterative_flow_blocks_without_a_renderer(tmp_path: Path) -> None:
    state = iterate({"surface": "report", "source": str(REPORT), "renders": str(RENDERS),
                     "workspace": str(tmp_path), "fixer_id": "executor", "max_rounds": 2})
    assert state["outcome"] == "blocked" and "renderer adapter" in state["reason"]
    assert (tmp_path / "iteration-history.json").is_file()


def test_review_template_tracks_current_policy_and_image_digest() -> None:
    state = snapshot("report", REPORT, RENDERS)
    draft = request("report", REPORT, RENDERS, "executor")
    assert draft["source_sha256"] == state["source_sha256"]
    assert draft["pages"][0]["image_sha256"] == state["pages"][0]["sha256"]
    assert draft["pages"][0]["observations"][0]["status"] == "pending"
    assert len(draft["pages"][0]["observations"][0]["criterion"]) > 25


def test_png_header_alone_cannot_pass_render_integrity(tmp_path: Path) -> None:
    from pbip_documenter.visual_quality.evidence import png_size

    screenshot = RENDERS / "01-executive-performance.png"
    truncated = tmp_path / "page.png"
    truncated.write_bytes(screenshot.read_bytes()[:24])
    import pytest

    with pytest.raises(ValueError, match="Incomplete|Truncated"):
        png_size(truncated)


def test_safe_geometry_fixer_is_dry_run_by_default(tmp_path: Path) -> None:
    from pbip_documenter.visual_quality.repair import align_kpi_rows

    clone = tmp_path / REPORT.name
    shutil.copytree(REPORT, clone)
    file = clone / "definition" / "pages" / "Executive_performance" / "visuals" / "card-096ca6e7" / "visual.json"
    value = json.loads(file.read_text(encoding="utf-8"))
    value["position"]["y"] -= 32
    file.write_text(json.dumps(value), encoding="utf-8")
    assert len(align_kpi_rows(clone)) == 1
    assert json.loads(file.read_text(encoding="utf-8"))["position"]["y"] == 104
    assert len(align_kpi_rows(clone, apply=True)) == 1
    assert json.loads(file.read_text(encoding="utf-8"))["position"]["y"] == 136
    assert not align_kpi_rows(clone)
    assert snapshot("report", clone, None)["static_pass"]


def test_remote_vision_endpoint_requires_explicit_permission(tmp_path: Path) -> None:
    import pytest

    from pbip_documenter.visual_quality.vision_reviewer import review_with_model

    with pytest.raises(ValueError, match="HTTPS and --allow-remote"):
        review_with_model(tmp_path / "request.json", RENDERS, tmp_path / "review.json",
                          endpoint="https://external.example/v1/chat/completions", model="vision",
                          reviewer_id="independent")


def test_local_vision_adapter_produces_actionable_rejections(tmp_path: Path) -> None:
    import io
    from unittest.mock import patch

    from pbip_documenter.visual_quality.vision_reviewer import review_with_model

    draft = request("report", REPORT, RENDERS, "fixer")
    request_file, response = tmp_path / "request.json", tmp_path / "review.json"
    request_file.write_text(json.dumps(draft), encoding="utf-8")
    observations = draft["pages"][0]["observations"]
    for row in observations:
        row.update(status="pass", reason="The reviewer specifically inspected this property on the rendered page.")
    observations[0].update(status="fail", severity="high", region=[0, 0, 1, 0.3],
                           proposed_fix="Increase the heading prominence.",
                           reason="The rendered header competes with surrounding filter boxes for attention.")
    wire = json.dumps({"choices": [{"message": {"content": json.dumps({"observations": observations})}}]})
    with patch("urllib.request.urlopen", side_effect=lambda *a, **k: io.BytesIO(wire.encode())) as mocked:
        result = review_with_model(request_file, RENDERS, response,
                                   endpoint="http://127.0.0.1:9999/v1/chat/completions",
                                   model="synthetic-test-vision", reviewer_id="independent")
    assert mocked.call_count == 5
    assert len(result["pages"]) == 5 and response.is_file()
    findings = audit("report", REPORT, RENDERS, response, "fixer")["findings"]
    assert sum(f["rule"] == "visual_defect" for f in findings) == 5


def test_iteration_repeats_render_review_and_repair_until_pass(tmp_path: Path) -> None:
    from unittest.mock import patch

    from pbip_documenter.visual_quality.runner import save

    cloned = tmp_path / REPORT.name
    renders, workspace = tmp_path / "renders", tmp_path / "run"
    shutil.copytree(REPORT, cloned)
    shutil.copytree(RENDERS, renders)
    calls = []

    def adapter(_spec: dict, kind: str, paths: dict[str, str]) -> None:
        calls.append(kind)
        if kind == "renderer":
            manifest = json.loads((renders / "capture-manifest.json").read_text(encoding="utf-8"))
            manifest["source_sha256"] = snapshot("report", cloned, None)["source_sha256"]
            (renders / "capture-manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
        elif kind == "reviewer":
            review = json.loads(Path(paths["request"]).read_text(encoding="utf-8"))
            review["reviewer"]["id"] = "separate-review-agent"
            for page in review["pages"]:
                for item in page["observations"]:
                    item.update(status="pass", reason="The rendered page was checked against this exact required criterion.")
            if calls.count("reviewer") == 1:
                issue = review["pages"][0]["observations"][0]
                issue.update(status="fail", severity="high", region=[0.1, 0.1, 0.5, 0.3],
                             reason="The title is visibly too small relative to the page's main metric.",
                             proposed_fix="Increase title prominence and check against the rendered page.")
            save(Path(paths["review"]), review)
        elif kind == "fixer":
            file = cloned / "definition" / "pages" / "Executive_performance" / "visuals" / "card-096ca6e7" / "visual.json"
            payload = json.loads(file.read_text(encoding="utf-8"))
            payload["position"]["x"] += 1
            file.write_text(json.dumps(payload), encoding="utf-8")

    spec = {"surface": "report", "source": str(cloned), "renders": str(renders),
            "workspace": str(workspace), "fixer_id": "fixer", "max_rounds": 3}
    with patch("pbip_documenter.visual_quality.runner._adapter", side_effect=adapter):
        result = iterate(spec)
    assert result["outcome"] == "passed" and len(result["rounds"]) == 2, result
    assert calls == ["renderer", "reviewer", "fixer", "renderer", "reviewer"]
