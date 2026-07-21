"""Document builder — assembles the full specification document."""

import os
import typing
from copy import deepcopy
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt

from pbip_documenter.analysis.connectors import _detect_connectors, _scan_all_sources
from pbip_documenter.analysis.dax import _dax_pattern, _transforms
from pbip_documenter.analysis.lineage import insert_lineage_diagram
from pbip_documenter.analysis.observations import generate_observations
from pbip_documenter.analysis.schema import insert_star_schema
from pbip_documenter.config import _SCHEMA_LANDSCAPE_MARGIN_CM, FONT, C
from pbip_documenter.docx_render.pagination import (
    _BLOCK_BUFFER_IN,
    _H3_IN,
    _H4_IN,
    estimate_table_block_h,
    estimate_wireframe_block_h,
    should_page_break,
    usable_page_height_in,
)
from pbip_documenter.docx_render.tables import (
    _OBS_CARD_COUNTER,
    _section_break_after_table,
    code_block,
    data_table,
    doc_control_table,
    kpi_strip,
    obs_card,
    prop_table,
)
from pbip_documenter.docx_render.template import _clear_template_body, _update_template_header
from pbip_documenter.docx_render.typography import (
    _remove_unnecessary_breaks,
    _run,
    body,
    bullet_item,
    callout,
    h1,
    h2,
    h3,
    h4,
    page_break,
    placeholder,
    suggested,
    suggested_bullet,
)
from pbip_documenter.docx_render.wireframe import insert_page_layout

# LANDSCAPE / PORTRAIT SECTION HANDLING
# Schema relationship diagrams are rendered on landscape pages to maximize width.
# Landscape pages intentionally use a minimal header with only a page number.
# Before entering landscape, the original portrait header/footer references are captured,
# then restored when returning to portrait. This avoids cloning image XML and prevents
# broken logo/header placeholders after the schema pages.
_SCHEMA_PORTRAIT_HEADER_FOOTER_REFS = None


def _clear_hdrftr(hdrftr: typing.Any) -> typing.Any:
    for child in list(hdrftr._element):
        hdrftr._element.remove(child)


def _capture_hdrftr_refs(section: typing.Any) -> typing.Any:
    refs: list[typing.Any] = []
    sect_pr = section._sectPr
    for tag in ("headerReference", "footerReference"):
        for node in sect_pr.findall(qn(f"w:{tag}")):
            refs.append(deepcopy(node))
    return refs


def _restore_hdrftr_refs(section: typing.Any, refs: typing.Any) -> typing.Any:
    sect_pr = section._sectPr
    for tag in ("headerReference", "footerReference"):
        for node in list(sect_pr.findall(qn(f"w:{tag}"))):
            sect_pr.remove(node)
    # Re-use the original header/footer parts instead of cloning header XML.
    # This keeps image/logo relationships valid and prevents broken-picture placeholders.
    for ref in reversed(refs or []):
        sect_pr.insert(0, deepcopy(ref))


def _add_page_number(paragraph: typing.Any) -> typing.Any:
    run = paragraph.add_run("Page ")
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_text)
    run._r.append(fld_end)
    run.font.name = FONT
    run.font.size = Pt(8)
    run.font.color.rgb = C.rgb(C.DGRAY)


def _set_landscape_page_number_header(section: typing.Any) -> typing.Any:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    _clear_hdrftr(section.header)
    _clear_hdrftr(section.footer)
    p = section.header.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    _add_page_number(p)


def _set_section_landscape(section: typing.Any, margin_cm: typing.Any = _SCHEMA_LANDSCAPE_MARGIN_CM) -> typing.Any:
    section.orientation = WD_ORIENT.LANDSCAPE
    if section.page_width < section.page_height:
        section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(margin_cm)
    section.bottom_margin = Cm(margin_cm)
    section.left_margin = Cm(margin_cm)
    section.right_margin = Cm(margin_cm)
    _set_landscape_page_number_header(section)
    return section


def _set_section_portrait(section: typing.Any) -> typing.Any:
    section.orientation = WD_ORIENT.PORTRAIT
    if section.page_width > section.page_height:
        section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.78)
    section.left_margin = Cm(1.78)
    section.right_margin = Cm(1.52)
    section.bottom_margin = Cm(1.52)
    # Unlink from previous section's header before restoring template refs,
    # otherwise the new section inherits the landscape page-number-only header.
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    _restore_hdrftr_refs(section, _SCHEMA_PORTRAIT_HEADER_FOOTER_REFS)
    return section


def _begin_schema_landscape_section(doc: typing.Any) -> typing.Any:
    global _SCHEMA_PORTRAIT_HEADER_FOOTER_REFS
    if _SCHEMA_PORTRAIT_HEADER_FOOTER_REFS is None:
        _SCHEMA_PORTRAIT_HEADER_FOOTER_REFS = _capture_hdrftr_refs(doc.sections[0])
    return _set_section_landscape(doc.add_section(WD_SECTION.NEW_PAGE))


def _end_schema_portrait_section(doc: typing.Any) -> typing.Any:
    return _set_section_portrait(doc.add_section(WD_SECTION.NEW_PAGE))


def _schema_has_relationship_diagrams(summary: typing.Any) -> typing.Any:
    sm = summary.get("semantic_model") or {}
    tables = {t.get("name") for t in (sm.get("tables", []) or []) if t.get("name")}
    return any(
        r.get("from_table") in tables and r.get("to_table") in tables for r in (sm.get("relationships", []) or [])
    )


# TWEAK: naming convention observation rules
# Flags low-signal object names for the Technical observations section.
# Add or remove terms in `exact` if local naming standards differ.
def _is_generic_name(name: typing.Any, kind: typing.Any = "object") -> typing.Any:
    norm = " ".join(str(name or "").strip().lower().replace("_", " ").replace("-", " ").split())
    exact: set[typing.Any] = {
        "table",
        "table1",
        "table 1",
        "new table",
        "query",
        "query1",
        "sheet1",
        "sheet 1",
        "column",
        "column1",
        "column 1",
        "field",
        "field1",
        "value",
        "values",
        "data",
        "measure",
        "measure1",
        "new measure",
        "calculation",
        "calc",
        "name",
        "description",
    }
    if norm in exact:
        return True
    if kind == "table" and norm.startswith(("table ", "query ", "sheet ")):
        return True
    return bool(kind in ("column", "measure") and norm.startswith(("column ", "field ", "measure ")))


def _naming_convention_observations(sm: typing.Any) -> typing.Any:
    tables = sm.get("tables", []) or []
    observations: list[typing.Any] = []
    generic_tables = [t.get("name", "") for t in tables if _is_generic_name(t.get("name"), "table")]
    generic_columns: list[typing.Any] = []
    generic_measures: list[typing.Any] = []
    table_prefixes: dict[typing.Any, typing.Any] = {"fact": 0, "dim": 0, "other": 0}
    inconsistent_examples: list[typing.Any] = []
    for t in tables:
        tname = str(t.get("name", "") or "")
        low = tname.lower().strip()
        if low.startswith(("fact", "fct")):
            table_prefixes["fact"] += 1
        elif low.startswith(("dim", "dimension")):
            table_prefixes["dim"] += 1
        else:
            table_prefixes["other"] += 1
            if len(inconsistent_examples) < 5:
                inconsistent_examples.append(tname)
        for c in t.get("columns", []) or []:
            if _is_generic_name(c.get("name", ""), "column"):
                generic_columns.append(f"{tname}.{c.get('name', '')}")
        for m in t.get("measures", []) or []:
            if _is_generic_name(m.get("name", ""), "measure"):
                generic_measures.append(f"{tname}.{m.get('name', '')}")
    if generic_tables:
        sample = ", ".join(generic_tables[:8]) + ("..." if len(generic_tables) > 8 else "")
        observations.append(
            (
                "Generic table names detected",
                f"Some tables use low-signal names ({sample}). Prefer business/domain names such as DimCustomer, FactSales, or ForecastSnapshot so lineage and support handover are easier to interpret.",
                C.MARIGOLD,
                "Warnings",
            )
        )
    if generic_columns:
        sample = ", ".join(generic_columns[:8]) + ("..." if len(generic_columns) > 8 else "")
        observations.append(
            (
                "Generic column names detected",
                f"Some columns use generic names ({sample}). Prefer descriptive business terms and avoid names like Column, Field, Value or Data unless they are intentionally generic staging artifacts.",
                C.MARIGOLD,
                "Warnings",
            )
        )
    if generic_measures:
        sample = ", ".join(generic_measures[:8]) + ("..." if len(generic_measures) > 8 else "")
        observations.append(
            (
                "Generic measure names detected",
                f"Some measures use low-signal names ({sample}). Prefer action-oriented names such as Net Sales, Units Sold, Forecast Accuracy, or YoY Growth.",
                C.MARIGOLD,
                "Warnings",
            )
        )
    total = sum(table_prefixes.values())
    if (
        total >= 8
        and table_prefixes["other"] >= max(3, total // 3)
        and (table_prefixes["fact"] or table_prefixes["dim"])
    ):
        sample = ", ".join([x for x in inconsistent_examples if x][:5])
        observations.append(
            (
                "Mixed table naming convention",
                f"The model mixes prefixed tables (Fact/Dim) with unprefixed tables. Examples without convention: {sample or 'n/a'}. Consider a consistent naming scheme for facts, dimensions, helpers, and snapshots.",
                C.SKY,
                "Info",
            )
        )
    if not generic_tables and not generic_columns and not generic_measures:
        observations.append(
            (
                "Naming conventions look serviceable",
                "No obviously generic table, column, or measure names were detected by the metadata scan. Continue to validate business terminology with report owners.",
                C.EVERGREEN,
                "Good Practices",
            )
        )
    return observations


def _render_generated_section(doc: typing.Any, title: typing.Any, section: typing.Any) -> typing.Any:
    h2(doc, title)
    if not section:
        suggested(doc, "No generated content available.")
        return

    body(doc, section.text)
    if section.assumptions:
        h3(doc, "Assumptions")
        for item in section.assumptions:
            suggested_bullet(doc, item)
    if section.missing_information:
        h3(doc, "Missing information to confirm")
        for item in section.missing_information:
            suggested_bullet(doc, item)
    if section.traceability_notes:
        h3(doc, "Traceability notes")
        for item in section.traceability_notes:
            bullet_item(doc, item)
    callout(doc, f"Generated source: {section.source}")


def _join_list(values: typing.Any, empty: typing.Any = "—") -> typing.Any:
    vals = [str(v) for v in (values or []) if v not in (None, "", "(not found)")]
    return ", ".join(vals) if vals else empty


def _render_semantic_model_notes(doc: typing.Any, sm: typing.Any) -> typing.Any:
    """Render concise semantic-model metadata for richer full-mode output."""
    parser_warnings = sm.get("parser_warnings") or sm.get("warnings") or []
    model = sm.get("model") or {}
    rows: list[typing.Any] = [
        ("Parser confidence", sm.get("parser_confidence", "unknown")),
        ("Warnings detected", str(len(parser_warnings))),
        ("Has RLS", "Yes" if sm.get("has_rls") else "No"),
        ("Hierarchies", str(sm.get("hierarchy_count", 0))),
        ("Calculation items", str(sm.get("calculation_item_count", 0))),
        ("Partitions", str(sm.get("partition_count", 0))),
    ]
    parser_mode = (sm.get("parser_mode") or "").strip()
    if parser_mode and parser_mode.lower() not in ("python_fallback", "fallback", "unknown"):
        rows.insert(0, ("Parser mode", parser_mode))

    h3(doc, "Model extraction notes")
    prop_table(doc, rows)

    h3(doc, "Model properties")
    prop_table(
        doc,
        [
            ("Culture", model.get("culture", "—")),
            ("Default data source version", model.get("data_source_version", "—")),
            ("Default mode", model.get("default_mode", "—")),
            ("Default data view", model.get("default_data_view", "—")),
            ("Query groups", _join_list(model.get("query_groups", []))),
            ("Model description", model.get("description", "—") or "—"),
        ],
    )

    if sm.get("data_sources"):
        h3(doc, "Detected data source endpoints")
        for src in sm.get("data_sources", []):
            bullet_item(doc, src)

    if parser_warnings:
        callout(
            doc,
            "Parsing warnings: "
            + "; ".join(str(w) for w in parser_warnings[:8])
            + ("; ..." if len(parser_warnings) > 8 else ""),
            kind="warn",
        )


def _render_hierarchies_and_calc_groups(doc: typing.Any, tables: typing.Any) -> typing.Any:
    rich_tables = [
        t for t in tables if t.get("hierarchies") or t.get("calculation_items") or t.get("is_calculation_group")
    ]
    if not rich_tables:
        return

    h3(doc, "Hierarchies and calculation groups")
    for t in rich_tables:
        p_h4 = h4(doc, f"[Table] {t.get('name', '?')}")
        p_h4.paragraph_format.keep_with_next = True

        if t.get("hierarchies"):
            h4(doc, "Hierarchies")
            rows: list[typing.Any] = []
            for h in t.get("hierarchies", []):
                levels = h.get("levels", []) or []
                level_names = [lvl.get("name", "?") for lvl in levels if lvl.get("name")]
                notes = ", ".join(level_names[:5]) if level_names else "Level detail was not extracted from metadata"
                rows.append([h.get("name", "?"), str(len(levels)), notes])
            data_table(
                doc, ["Hierarchy", "Levels detected", "What was identified"], rows, [2400, 1100, 5034], compact=True
            )

        if t.get("is_calculation_group") or t.get("calculation_items"):
            h4(doc, "Calculation groups")
            items = t.get("calculation_items", []) or []
            item_names = ", ".join((item.get("name") or "?") for item in items[:8]) if items else "—"
            data_table(
                doc,
                ["Table", "Items detected", "Item names (sample)"],
                [[t.get("name") or "?", str(len(items)), item_names]],
                [2400, 1100, 5034],
                compact=True,
            )


def _render_relationship_details(doc: typing.Any, rels: typing.Any) -> typing.Any:
    if not rels:
        return
    has_cardinality = any((r.get("cardinality") or "").strip() not in ("", "—") for r in rels)
    has_cross = any((r.get("cross_filtering") or "").strip() not in ("", "—", "single") for r in rels)
    has_security = any((r.get("security_filtering") or "").strip() not in ("", "—") for r in rels)
    has_join = any((r.get("join_on_date_behavior") or "").strip() not in ("", "—") for r in rels)
    if not any([has_cardinality, has_cross, has_security, has_join]):
        return

    headers: list[typing.Any] = ["Relationship"]
    widths: list[typing.Any] = [4200]
    if has_cardinality:
        headers.append("Cardinality")
        widths.append(1200)
    if has_cross:
        headers.append("Cross-filter")
        widths.append(1200)
    if has_security:
        headers.append("Security filtering")
        widths.append(1400)
    if has_join:
        headers.append("Join on date")
        widths.append(1400)
    headers.append("Active")
    widths.append(800)

    rows: list[typing.Any] = []
    for r in rels:
        row: list[typing.Any] = [
            f"{r.get('from_table', '?')}.{r.get('from_field', '?')} → {r.get('to_table', '?')}.{r.get('to_field', '?')}"
        ]
        if has_cardinality:
            row.append(r.get("cardinality", "—"))
        if has_cross:
            row.append(r.get("cross_filtering", "—"))
        if has_security:
            row.append(r.get("security_filtering", "—"))
        if has_join:
            row.append(r.get("join_on_date_behavior", "—"))
        row.append("Yes" if r.get("is_active", True) else "No")
        rows.append(row)

    h3(doc, "Relationship details")
    data_table(doc, headers, rows, widths, compact=True, repeat_header=True)


def _render_rls_details(doc: typing.Any, roles: typing.Any) -> typing.Any:
    if not roles:
        return

    h3(doc, "Role detail and filter expressions")
    for idx, role in enumerate(roles, start=1):
        role_name = (
            (role.get("name") or "").strip() or Path(role.get("file", "")).stem.replace("_", " ") or f"Role {idx}"
        )
        badge = " [personal-name-like]" if role.get("looks_like_personal_name") else ""
        p_h4 = h4(doc, f"[Role] {role_name}{badge}")
        p_h4.paragraph_format.keep_with_next = True
        prop_table(
            doc,
            [
                ("Model permission", role.get("model_permission", "read")),
                ("Filtered tables", str(role.get("filtered_table_count", len(role.get("table_permissions", []))))),
                ("Complexity", role.get("role_filter_complexity", "—")),
                ("Source file", role.get("file", "—")),
            ],
        )

        perms = role.get("table_permissions", []) or []
        if perms:
            preview_rows: list[typing.Any] = []
            for tp in perms:
                expr_preview = (tp.get("expression", "—") or "—").replace("\n", " ")
                if len(expr_preview) > 120:
                    expr_preview = expr_preview[:120] + "..."
                preview_rows.append([tp.get("table", "?"), expr_preview])
            data_table(doc, ["Table", "Filter expression (preview)"], preview_rows, [2200, 7834], compact=True)
            for tp in perms:
                h4(doc, f"[RLS Filter] {tp.get('table', '?')}")
                code_block(doc, tp.get("expression", "—"))
        else:
            body(doc, "No table-level filters defined for this role.")


def _render_user_defined_functions(doc: typing.Any, functions: typing.Any, mode: typing.Any) -> typing.Any:
    if not functions:
        return
    h3(doc, "User-defined DAX functions")
    func_rows: list[typing.Any] = []
    for f in functions:
        func_rows.append(
            [
                f["name"],
                f.get("params", ""),
                f.get("return_type", ""),
                "\n".join(f.get("documentation", [])) or "\u2014",
            ]
        )
    data_table(
        doc,
        ["Function Name", "Parameters", "Return Type", "Documentation"],
        func_rows,
        [2800, 2200, 1600, 3034],
        compact=True,
        repeat_header=True,
        title="User-defined functions",
    )
    if mode == "full":
        for f in functions:
            _h4 = h4(
                doc,
                f"[UDF]  {f['name']}({f.get('params', '')})  \u2192  {f.get('return_type', '')}",
            )
            _h4.paragraph_format.keep_with_next = True
            code_block(doc, f.get("body", "\u2014 no body \u2014"))
    else:
        callout(doc, f"Full DAX for all {len(functions)} user-defined functions available in Full mode.")


def _render_visual_actions(
    doc: typing.Any, visuals: typing.Any, pages_map: typing.Any, bookmarks_map: typing.Any, mode: typing.Any
) -> typing.Any:
    """Render visual actions / links for a single page."""
    actions = [v for v in visuals if v.get("action")]
    if not actions:
        return

    rows: list[typing.Any] = []
    for v in actions:
        act = v["action"]
        act_type = act.get("type", "?")
        visual_label = v.get("button_text", "") or v.get("visual_type", "?")
        v.get("id", "")[:8]

        target = "\u2014"
        detail = "\u2014"

        if act_type == "PageNavigation":
            page_id = act.get("target_page_id", "")
            target = pages_map.get(page_id, page_id) if page_id else "\u2014"
            detail = f"Navigate to page: {target}"
        elif act_type == "Bookmark":
            bm_id = act.get("target_bookmark_id", "")
            target = bookmarks_map.get(bm_id, bm_id) if bm_id else "\u2014"
            detail = f"Apply bookmark: {target}"
        elif act_type == "DataFunction":
            fn = act.get("function_name", "?")
            target = fn
            params = act.get("parameters", [])
            if params:
                param_summary = "; ".join(f"{p['name']}={p.get('value', '?')}" for p in params)
                detail = f"DataFunction: {fn}({param_summary})"
            else:
                detail = f"DataFunction: {fn}()"
        elif act_type == "ClearAllSlicers":
            target = "All slicers"
            detail = "Reset all slicers on this page"
        elif act_type == "Back":
            target = "Previous page"
            detail = act.get("tooltip", "Return to previous page")
        else:
            detail = act.get("tooltip", "")

        rows.append([visual_label, act_type, target, detail])

    p_h4 = h4(doc, "Actions & Navigation")
    p_h4.paragraph_format.keep_with_next = True
    data_table(doc, ["Visual", "Action Type", "Target", "Details"], rows, [2000, 1800, 2000, 4234], compact=True)

    if mode == "full":
        for v in actions:
            act = v["action"]
            if act.get("type") == "DataFunction" and act.get("parameters"):
                fn = act.get("function_name", "?")
                p_h4_df = h4(doc, f"[DataFunction] {fn}")
                p_h4_df.paragraph_format.keep_with_next = True
                param_rows: list[typing.Any] = []
                for p in act["parameters"]:
                    param_rows.append(
                        [
                            p.get("name", "?"),
                            p.get("data_type", "?"),
                            "Yes" if p.get("is_optional") else "No",
                            p.get("value", "\u2014"),
                        ]
                    )
                data_table(
                    doc,
                    ["Parameter", "Data Type", "Optional", "Value Binding"],
                    param_rows,
                    [2000, 1600, 1200, 5234],
                    compact=True,
                )


def _render_bookmarks_section(
    doc: typing.Any,
    bookmarks: typing.Any,
    mode: typing.Any,
    pages_map: typing.Any = None,
    visuals_map: typing.Any = None,
) -> typing.Any:
    """Render dedicated bookmarks section."""
    if not bookmarks:
        return

    h2(doc, "3.4.1 Bookmarks")

    if mode != "full":
        body(
            doc,
            f"{len(bookmarks)} bookmark(s) configured. The table below provides an overview of each bookmark's target visuals, filters, and triggers.",
        )
        rows: list[typing.Any] = []
        has_groups = any(bm.get("group") for bm in bookmarks)
        for bm in bookmarks:
            name = bm.get("display_name", bm.get("id", "?"))
            group = bm.get("group", "\u2014")

            # Standard mode: summarise target count + page instead of raw IDs
            tv = bm.get("target_visuals", [])
            if tv:
                n = len(tv)
                page_id = bm.get("active_section", "")
                page_name = pages_map.get(page_id, page_id) if pages_map else page_id
                page_label = f"{page_name} ({page_id})" if page_name and page_name != page_id else page_id or "—"
                targets = f"{n} visual(s) on {page_label}" if page_id else f"{n} visual(s)"
            else:
                targets = "\u2014"

            filters = ", ".join(bm.get("filter_entities", [])) or "\u2014"
            triggered = bm.get("triggered_by", [])
            if triggered:
                trigger_summary = "; ".join(
                    f"{t.get('button_text', t.get('visual_type', '?'))} on {t.get('page_name', '?')}"
                    for t in triggered[:3]
                )
                if len(triggered) > 3:
                    trigger_summary += f" (+{len(triggered) - 3} more)"
            else:
                trigger_summary = "\u2014"
            row = (
                [name, group, targets, filters, trigger_summary]
                if has_groups
                else [name, targets, filters, trigger_summary]
            )
            rows.append(row)

        headers: list[typing.Any]
        widths: list[typing.Any]
        if has_groups:
            headers = ["Bookmark", "Group", "Target Visuals", "Filter Entities", "Triggered By"]
            widths = [2000, 1600, 2400, 2000, 2034]
        else:
            headers = ["Bookmark", "Target Visuals", "Filter Entities", "Triggered By"]
            widths = [2400, 2800, 2400, 2434]
        data_table(doc, headers, rows, widths, compact=True, repeat_header=True, title="Bookmarks")
        return

    # Full mode: detailed per-bookmark blocks
    from collections import Counter, defaultdict

    for bm in bookmarks:
        name = bm.get("display_name", bm.get("id", "?"))
        p_h3 = h3(doc, f"[Bookmark] {name}")
        p_h3.paragraph_format.keep_with_next = True

        # Metadata
        details: list[typing.Any] = []
        if bm.get("active_section"):
            active_id = bm["active_section"]
            active_name = pages_map.get(active_id, active_id) if pages_map else active_id
            active_label = f"{active_name} ({active_id})" if active_name and active_name != active_id else active_id
            details.append(("Active section", active_label))
        if bm.get("filter_entities"):
            details.append(("Filter entities", ", ".join(bm["filter_entities"])))
        if bm.get("triggered_by"):
            triggers = "; ".join(
                f"{t.get('button_text', t.get('visual_type', '?'))} ({t.get('visual_id', '?')[:8]}) on {t.get('page_name', '?')}"
                for t in bm["triggered_by"]
            )
            details.append(("Triggered by", triggers))

        if details:
            prop_table(doc, details)
        else:
            body(doc, "No additional details available.")

        # Target visuals
        tv = bm.get("target_visuals", [])
        if tv:
            counts: Counter[str] = Counter()
            named: list[typing.Any] = []
            ids_by_type = defaultdict(list)
            for v_id in tv:
                vis_info = visuals_map.get(v_id) if visuals_map else None
                v_type = vis_info.get("visual_type", "") if vis_info else ""
                if not v_type or v_type == "not found" or v_type == "(not found)":
                    v_type = "Other"
                counts[v_type] += 1
                ids_by_type[v_type].append(v_id)
                btn = vis_info.get("button_text", "") if vis_info else ""
                if btn:
                    named.append(f"{btn} ({v_type})")

            body(doc, "")

            p_h4_targets = h4(doc, "Target visuals")
            p_h4_targets.paragraph_format.keep_with_next = True

            if counts:
                summary_rows: list[typing.Any] = []
                for vtype, cnt in counts.most_common():
                    ids_str = ", ".join(ids_by_type[vtype])
                    summary_rows.append([vtype, str(cnt), ids_str])
                data_table(doc, ["Visual Type", "Count", "Visual IDs"], summary_rows, [2400, 1200, 5434], compact=True)

            # Gap between count table and named visuals
            body(doc, "")

            if named:
                p_h4_named = h4(doc, "Named visuals")
                p_h4_named.paragraph_format.keep_with_next = True
                for item in named:
                    bullet_item(doc, item)
        else:
            body(doc, "No target visuals configured.")

        # Spacer between bookmark blocks
        body(doc, "")


def _render_custom_visuals_section(doc: typing.Any, custom_visuals: typing.Any) -> typing.Any:
    """Render dedicated custom visuals section."""
    if not custom_visuals:
        return

    h2(doc, "3.4.2 Custom Visuals")
    body(
        doc,
        "The following third-party custom visuals are used in this report. These are not part of the standard Power BI visual library and may require additional licensing or maintenance.",
    )

    for cv in custom_visuals:
        name = cv
        if len(cv) > 32:
            for prefix in ["upSetJSVenn", "textSearchSlicer"]:
                if cv.startswith(prefix):
                    name = f"{prefix} ({cv})"
                    break
        bullet_item(doc, name)


def _collect_data_functions(pages: typing.Any) -> typing.Any:
    """Aggregate all DataFunction actions across report pages."""
    from collections import defaultdict

    func_map: defaultdict[str, dict[str, typing.Any]] = defaultdict(
        lambda: {
            "function_name": "",
            "auto_refresh": False,
            "usages": 0,
            "pages": set(),
            "visuals": [],
            "parameters": [],
        }
    )

    for pg in pages:
        pg_name = pg.get("display_name", pg.get("id", "?"))
        for v in pg.get("visuals", []):
            act = v.get("action")
            if not act or act.get("type") != "DataFunction":
                continue
            fn = act.get("function_name", "?")
            if not fn or fn == "?":
                continue
            entry = func_map[fn]
            entry["function_name"] = fn
            entry["auto_refresh"] = act.get("auto_refresh", False)
            entry["usages"] += 1
            entry["pages"].add(pg_name)
            entry["visuals"].append(
                {
                    "page": pg_name,
                    "visual_type": v.get("visual_type", "?"),
                    "button_text": v.get("button_text", "") or v.get("visual_type", "?"),
                    "id": v.get("id", "")[:8],
                }
            )
            if not entry["parameters"] and act.get("parameters"):
                entry["parameters"] = act["parameters"]

    result: list[typing.Any] = []
    for fn in sorted(func_map.keys()):
        entry = func_map[fn]
        entry["pages"] = sorted(entry["pages"])
        result.append(entry)
    return result


def _render_data_functions_section(doc: typing.Any, data_functions: typing.Any, mode: typing.Any) -> typing.Any:
    """Render dedicated Data Functions section."""
    if not data_functions:
        return

    h2(doc, "3.4.3 Data Functions")
    body(
        doc,
        "The following data functions are invoked by report visuals. These functions typically perform write-back, alerting, or other interactive operations against the data source.",
    )

    rows: list[typing.Any] = []
    for df in data_functions:
        rows.append(
            [
                df["function_name"],
                str(df["usages"]),
                "Yes" if df["auto_refresh"] else "No",
                ", ".join(df["pages"]),
                str(len(df["parameters"])),
            ]
        )
    data_table(
        doc,
        ["Function Name", "Usages", "Auto-Refresh", "Pages Used", "Parameters"],
        rows,
        [2800, 1000, 1400, 3034, 1800],
        compact=True,
        repeat_header=True,
        title="Data functions",
    )

    callout(
        doc,
        "Data functions perform write-back or alerting operations. Review security, audit, and data-integrity implications before deployment.",
        kind="warn",
    )

    if mode == "full":
        for df in data_functions:
            fn = df["function_name"]
            p_h4 = h4(doc, f"[DataFunction]  {fn}")
            p_h4.paragraph_format.keep_with_next = True

            prop_table(
                doc,
                [
                    ("Function name", fn),
                    ("Auto-refresh", "Yes" if df["auto_refresh"] else "No"),
                    ("Total usages", str(df["usages"])),
                    ("Pages used", ", ".join(df["pages"])),
                ],
            )

            if df["parameters"]:
                param_rows: list[typing.Any] = []
                for p in df["parameters"]:
                    param_rows.append(
                        [
                            p.get("name", "?"),
                            p.get("data_type", "?"),
                            "Yes" if p.get("is_optional") else "No",
                            p.get("value", "\u2014"),
                        ]
                    )
                data_table(
                    doc,
                    ["Parameter", "Data Type", "Optional", "Value Binding"],
                    param_rows,
                    [2000, 1600, 1200, 5234],
                    compact=True,
                    title=f"Parameters for {fn}",
                )
            else:
                body(doc, "No parameters defined.")

            usage_rows: list[typing.Any] = []
            for v in df["visuals"]:
                usage_rows.append(
                    [
                        v["page"],
                        v["visual_type"],
                        v["button_text"],
                        v["id"],
                    ]
                )
            data_table(
                doc,
                ["Page", "Visual Type", "Button Text", "Visual ID"],
                usage_rows,
                [2800, 1600, 3034, 2600],
                compact=True,
                title=f"Usages of {fn}",
            )


def build_doc(
    summary: typing.Any,
    mode: typing.Any = "default",
    logo_path: typing.Any = None,
    template_path: typing.Any = None,
    augmentation: typing.Any = None,
) -> typing.Any:
    LOCAL_ONLY = os.environ.get("PBIP_DOCUMENTER_LOCAL_ONLY", "1").lower() not in ("0", "false", "no")
    if LOCAL_ONLY:
        augmentation = None
    sm = summary.get("semantic_model") or {}
    rpt = summary.get("report") or {}
    tables = sm.get("tables", [])
    rels = sm.get("relationships", [])
    exprs = sm.get("expressions", [])
    pages = rpt.get("pages", [])
    report_name = rpt.get("display_name", summary.get("project_name", ""))
    culture = sm.get("model", {}).get("culture", "\u2014")
    compatibility_level = sm.get("compatibility_level", "\u2014")

    if template_path and os.path.isfile(template_path):
        print(f"  Using template: {template_path}")
        doc = Document(template_path)
        _clear_template_body(doc)
        _update_template_header(doc, report_name, rpt.get("logical_id", ""), sm.get("logical_id", ""))
        # Capture template header/footer refs immediately while intact —
        # before any section breaks overwrite them.
        global _SCHEMA_PORTRAIT_HEADER_FOOTER_REFS
        _SCHEMA_PORTRAIT_HEADER_FOOTER_REFS = _capture_hdrftr_refs(doc.sections[0])
    else:
        doc = Document()
        sec = doc.sections[0]
        sec.page_width = Emu(11906 * 635)
        sec.page_height = Emu(16838 * 635)
        sec.top_margin = Cm(1.78)
        sec.left_margin = Cm(1.78)
        sec.right_margin = Cm(1.52)
        sec.bottom_margin = Cm(1.52)
        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hp2 = sec.header.add_paragraph()
        hp2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _run(hp2, f"{report_name} \u2014 Design Specification", size=8, italic=True, color=C.DGRAY)
        h1(doc, "Table of Contents")
        callout(doc, "Right-click \u2192 Update Field to refresh page numbers after opening in Word.")
        page_break(doc)

    h1(doc, "Document control")
    doc_control_table(
        doc,
        [
            ("Solution / Report Name", report_name),
            ("Project / CMDB ID", rpt.get("logical_id", "\u2014")),
            ("Version", "1.0"),
            ("Document Date", datetime.now().strftime("%Y-%m-%d")),
            ("Document Owner", "[To be confirmed]"),
            ("Business Owner", "[To be confirmed]"),
            ("Classification / GxP", "[Proprietary] | Non-GxP"),
            ("Related Change / Release", "[To be confirmed]"),
        ],
    )

    h1(doc, "1. Purpose and scope")
    callout(
        doc,
        "Content below was inferred from PBIP metadata. Items marked [SUGGESTED] in pink should be reviewed and confirmed by the business owner.",
    )
    if LOCAL_ONLY:
        callout(doc, "External Jira / Power BI enrichments are disabled in local-only mode.")
    elif augmentation and augmentation.warning:
        callout(doc, augmentation.warning)
    h2(doc, "1.1 Purpose")
    if augmentation and augmentation.requirement:
        body(doc, augmentation.requirement.text)
    else:
        suggested(doc, "Describe the business problem this report solves and the intended outcome.")
    h2(doc, "1.2 Users and usage")
    suggested(doc, "List user groups, the pages they use, and decisions they make.")
    h2(doc, "1.3 In scope")
    total_m = sum(len(t.get("measures", [])) for t in tables)
    suggested(
        doc,
        f"{len(pages)} pages \u2014 {', '.join(p.get('display_name', '?') for p in pages)}. "
        f"{len(tables)} tables, {total_m} measures. Data from {len(set(sm.get('data_sources', [])))} source(s).",
    )
    h2(doc, "1.4 Out of scope")
    for item in [
        "No mobile-optimized layout",
        "No real-time or streaming refresh",
        "No data write-back or workflow automation",
        "No budget/forecast comparisons",
        "No historical data prior to source inception",
    ]:
        suggested_bullet(doc, item)
    h2(doc, "1.5 Assumptions, dependencies and limitations")
    if augmentation and augmentation.requirement and augmentation.requirement.assumptions:
        for item in augmentation.requirement.assumptions:
            suggested_bullet(doc, item)
        for item in augmentation.requirement.missing_information:
            suggested_bullet(doc, item)
    else:
        suggested(doc, "List assumptions, dependencies (e.g. SharePoint sites, mapping files), and limitations.")
    page_break(doc)

    h1(doc, "2. Solution and operations summary")
    h2(doc, "2.1 Delivered artifacts")
    # Use augmentation data if available, otherwise fall back to placeholders
    workspace_display: typing.Any = "[To be confirmed]"
    app_display: typing.Any = report_name
    jira_links: list[typing.Any] = []

    if augmentation and augmentation.report_context:
        ctx = augmentation.report_context

        # Build workspace display with link - using tuple (text, url) for hyperlinks
        if ctx.get("workspace_name"):
            workspace_name = ctx.get("workspace_name")
            workspace_url = ctx.get("workspace_url") or ctx.get("web_url")
            workspace_display = [(workspace_name, workspace_url)] if workspace_url else workspace_name

        # Build app/report display with link - using tuple (text, url) for hyperlinks
        if ctx.get("app_name"):
            app_name = ctx.get("app_name")
            app_url = ctx.get("app_url")
            report_url = ctx.get("web_url")
            if app_url:
                app_display = [f"{app_name} / {report_name}", ("Open in Power BI App", app_url)]
            elif report_url:
                app_display = [f"{app_name} / {report_name}", ("Open Report", report_url)]
            else:
                app_display = f"{app_name} / {report_name}"
        elif ctx.get("web_url"):
            web_url = ctx.get("web_url")
            app_display = [report_name, ("Open Report", web_url)]

        # Build Jira links from matched issues - using tuple (text, url) for hyperlinks
        matched_jira = ctx.get("matched_jira", [])
        if matched_jira:
            for jira_match in matched_jira[:3]:  # Top 3 matches
                issue_key = jira_match.get("issue_key")
                jira_url = jira_match.get("jira_url")
                jira_summary = jira_match.get("summary", "")
                if jira_url and issue_key:
                    # Display as: KEY-123: Summary text (clickable)
                    display_text = f"{issue_key}: {jira_summary}" if jira_summary else issue_key
                    jira_links.append((display_text, jira_url))
                elif issue_key:
                    jira_links.append(f"{issue_key}: {jira_summary}" if jira_summary else issue_key)

    # Build artifacts table rows
    artifact_rows: list[typing.Any] = [
        ("Workspace(s)", workspace_display),
        ("App / Report(s)", app_display),
        ("Semantic model(s)", sm.get("display_name", "\u2014")),
        (
            "Related assets",
            ", ".join(
                filter(None, [rpt.get("themes", {}).get("custom_theme"), rpt.get("themes", {}).get("base_theme")])
            ),
        ),
    ]

    # Add Jira references if available
    if jira_links:
        artifact_rows.append(("Jira References", jira_links))

    artifact_rows.append(("Deployment path", "[To be confirmed]"))

    prop_table(doc, artifact_rows)
    h2(doc, "2.2 Connectivity and refresh")
    prop_table(
        doc,
        [
            ("Connectivity mode", rpt.get("dataset_mode", "Import")),
            ("Refresh schedule", "[To be confirmed]"),
            ("Refresh owner / monitor", "[To be confirmed]"),
            ("Gateway / special connectivity", "[To be confirmed]"),
            (
                "Operational notes",
                f"Culture: {culture}, Compatibility level: {compatibility_level}",
            ),
        ],
    )
    h2(doc, "2.3 Access and security summary")
    suggested_bullet(
        doc,
        "Development and testing in a dedicated Power BI workspace. Published via Power BI App with audience-based access.",
    )
    roles_exist = sm.get("roles") and len(sm.get("roles", [])) > 0
    suggested_bullet(
        doc,
        "RLS roles present \u2014 see Section 3.5."
        if roles_exist
        else "No RLS roles defined \u2014 all app users see the same dataset.",
    )
    h2(doc, "2.4 References")
    placeholder(doc, "Backlog / demand record")
    placeholder(doc, "Support runbook / access document")
    page_break(doc)

    h1(doc, "3. Technical specifications")
    callout(
        doc, f"Mode: {'Full' if mode == 'full' else 'Default (concise)'}. All statements derived from PBIP metadata."
    )
    if LOCAL_ONLY:
        callout(doc, "Running in local-only mode: cache-backed Jira and Power BI service augmentations are disabled.")
    elif augmentation:
        callout(
            doc, "Technical narrative may include cache-backed generated context and must be reviewed before approval."
        )
    insert_lineage_diagram(doc, summary)
    page_break(doc)

    h2(doc, "3.1 Semantic model overview")
    total_cols = sum(len(t.get("columns", [])) for t in tables)
    total_meas = sum(len(t.get("measures", [])) for t in tables)
    kpi_strip(
        doc,
        [
            {"label": "Tables", "value": str(len(tables)), "color": C.PACIFIC},
            {"label": "Columns", "value": str(total_cols), "color": C.PACIFIC},
            {"label": "Measures", "value": str(total_meas), "color": C.RUBINE},
            {"label": "Relationships", "value": str(len(rels)), "color": C.EVERGREEN},
            {"label": "Pages", "value": str(len(pages)), "color": C.SKY},
            {"label": "Visuals", "value": str(sum(len(p.get("visuals", [])) for p in pages)), "color": C.SKY},
        ],
    )
    if _schema_has_relationship_diagrams(summary):
        _begin_schema_landscape_section(doc)
        # TWEAK: schema section heading style
        # This is the section-level heading for the whole schema diagram group.
        # The per-diagram title is drawn inside each DrawingML canvas in analysis/schema.py,
        # which prevents title-only orphan pages.
        # To adjust section heading size, change Pt(13) below. Keep >= 8 pt.
        p_schema_h = h3(doc, "Semantic model relationship diagrams")
        p_schema_h.paragraph_format.keep_with_next = True
        for _r in p_schema_h.runs:
            _r.font.size = Pt(13)
            _r.font.bold = True
        pattern = insert_star_schema(doc, summary)
        _end_schema_portrait_section(doc)
    else:
        pattern = insert_star_schema(doc, summary)
    body(
        doc,
        (
            f'Model: "{sm.get("display_name", report_name)}". {rpt.get("dataset_mode", "Import")} mode, '
            f"compatibility level {compatibility_level}. "
            f"{len(tables)} tables \u00b7 {total_cols} columns \u00b7 {total_meas} measures \u00b7 "
            f"{len(rels)} relationships. Pattern: {pattern}."
        ),
    )
    if mode == "full":
        _render_semantic_model_notes(doc, sm)

    h3(doc, "3.1.1 Table inventory")
    trows: list[typing.Any] = []
    for t in tables:
        n = t.get("name") or "?"
        role = (
            "Fact"
            if any(n.lower().startswith(p) for p in ("fact_", "fct_"))
            else "Dimension"
            if any(n.lower().startswith(p) for p in ("dim_", "dimension_"))
            else "Other"
        )
        trows.append(
            [
                n,
                role,
                "Import",
                str(len(t.get("columns", []))),
                str(len(t.get("measures", []))),
                "Yes" if t.get("is_hidden") else "",
                t.get("data_category", ""),
            ]
        )
    data_table(
        doc,
        ["Table", "Role", "Storage Mode", "Columns", "Measures", "Hidden", "Data Category"],
        trows,
        [2600, 1200, 900, 700, 700, 800, 1000],
        compact=True,
        repeat_header=True,
        title="Table inventory",
    )
    if mode == "full":
        h3(doc, "3.1.2 Column detail")
        for t in tables:
            p_col_h4 = h4(doc, f"[Table]  {t.get('name', '?')}{'  [HIDDEN]' if t.get('is_hidden') else ''}")
            p_col_h4.paragraph_format.keep_with_next = True
            cr = [
                [
                    c.get("name", ""),
                    c.get("data_type", "?"),
                    c.get("summarize_by", ""),
                    c.get("sort_by_column", ""),
                    c.get("format_string", ""),
                    "Y" if c.get("is_hidden") else "",
                ]
                for c in t.get("columns", [])
            ]
            if cr:
                data_table(
                    doc,
                    ["Column", "Data Type", "Summarize By", "Sort By Column", "Format String", "Hidden"],
                    cr,
                    compact=True,
                )
        _render_hierarchies_and_calc_groups(doc, tables)
    h3(doc, "3.1.3 Relationships" if mode == "full" else "Relationships")
    rrows = [
        [
            str(i + 1),
            r.get("from_table", "?"),
            r.get("from_field", "?"),
            r.get("to_table", "?"),
            r.get("to_field", "?"),
            r.get("cross_filtering", "single"),
            "Yes" if r.get("is_active", True) else "No",
        ]
        for i, r in enumerate(rels)
    ]
    data_table(
        doc,
        ["#", "From Table", "From Column", "To Table", "To Column", "Direction", "Active"],
        rrows,
        [400, 1800, 1500, 1800, 1500, 900, 800],
        compact=True,
        repeat_header=True,
        title="Relationships",
    )
    if mode == "full":
        _render_relationship_details(doc, rels)
    page_break(doc)

    h2(doc, "3.2 Data sources and Power Query overview")
    h3(doc, "Source systems")
    all_sources = _scan_all_sources(exprs, tables)
    srows = [
        [
            src["fn"].split(".")[0],
            f"{src['label']} ({src['cat']})",
            src["server"] if src["server"] != "(auto-detect)" else "\u2014",
            f"{src['table_count']} table(s)",
        ]
        for src in all_sources
    ] or [["(none detected)", "\u2014", "No data source calls found", "\u2014"]]
    data_table(doc, ["Connector", "Type", "Server / URL", "Tables"], srows, [1600, 2000, 5234, 1200], compact=True)
    h3(doc, "Query inventory")
    qrows = [
        [
            t.get("name") or "?",
            "Calc Table" if p.get("type") == "calculated" else "Partition",
            "Yes",
            p.get("mode", "?"),
            _transforms(p.get("m_expression", "")),
        ]
        for t in tables
        for p in t.get("partitions", [])
    ]
    qrows += [[e.get("name") or "?", f"Expr ({e.get('result_type', '?')})", "No", "\u2014", ""] for e in exprs]
    data_table(
        doc,
        ["Query / Expression", "Category", "Loaded", "Storage Mode", "Transforms Applied"],
        qrows,
        [2600, 1400, 800, 800, 4434],
        compact=True,
        repeat_header=True,
        title="Query inventory",
    )
    if mode == "full":
        _section_break_after_table(doc, len(qrows))
        shared_exprs = [e for e in exprs if e.get("name")]
        if shared_exprs:
            h3(doc, "Shared expressions and parameters")
            erows: list[typing.Any] = []
            for e in shared_exprs:
                preview = (e.get("expression", "") or "—").replace("\n", " ")
                if len(preview) > 90:
                    preview = preview[:87] + "..."
                erows.append([e.get("name", "?"), e.get("result_type", "—"), e.get("query_group", "—"), preview])
            data_table(
                doc,
                ["Expression", "Result Type", "Query Group", "Preview"],
                erows,
                [2200, 1200, 1600, 5034],
                compact=True,
                repeat_header=True,
            )
        h3(doc, "Full M / Power Query code")
        for t in tables:
            for p in t.get("partitions", []):
                mc = p.get("m_expression", "")
                cd = p.get("calculated_dax", "")
                if mc:
                    _h4 = h4(doc, f"[Table]  {t['name']}  ({p.get('mode', '?')})")
                    _h4.paragraph_format.keep_with_next = True
                    code_block(doc, mc)
                elif cd:
                    _h4 = h4(doc, f"[Table]  {t['name']}  (calculated)")
                    _h4.paragraph_format.keep_with_next = True
                    code_block(doc, cd)
        for e in exprs:
            if e.get("expression"):
                _h4 = h4(doc, f"[Expression]  {e['name']}  ({e.get('result_type', '?')})")
                _h4.paragraph_format.keep_with_next = True
                code_block(doc, e["expression"])
    else:
        callout(doc, f"Full M code for {len(exprs)} expressions and {len(tables)} partitions available in Full mode.")
    page_break(doc)

    h2(doc, "3.3 Measures and calculations")
    all_m = [{**m, "table": t.get("name") or "?"} for t in tables for m in t.get("measures", [])]
    mrows = [
        [str(i + 1), m["name"], m["table"], m.get("format_string", "\u2014"), _dax_pattern(m.get("expression", ""))]
        for i, m in enumerate(all_m)
    ]
    data_table(
        doc,
        ["#", "Measure Name", "Source Table", "Format String", "DAX Pattern"],
        mrows,
        [400, 2800, 2200, 1600, 3034],
        compact=True,
        repeat_header=True,
        title="Measures",
    )
    if mode == "full":
        _section_break_after_table(doc, len(mrows))
        h3(doc, "Full DAX code")
        for m in all_m:
            _h4 = h4(
                doc,
                f"[{m['table']}]  {m['name']}{('  [' + m['format_string'] + ']') if m.get('format_string') else ''}",
            )
            _h4.paragraph_format.keep_with_next = True
            code_block(doc, m.get("expression", "\u2014 no expression \u2014"))
    else:
        callout(doc, f"Full DAX for all {len(all_m)} measures available in Full mode.")
    page_break(doc)

    # 3.3.1 User-defined DAX functions
    functions = sm.get("functions", [])
    _render_user_defined_functions(doc, functions, mode)

    h2(doc, "3.4 Report structure")
    pgrows: list[typing.Any] = []
    for pg in pages:
        vts = pg.get("visual_type_summary", {})
        tot = pg.get("visual_count", len(pg.get("visuals", [])))
        dv = sum(v for k, v in vts.items() if k not in ("shape", "basicShape", "image", "textbox"))
        sl = vts.get("slicer", 0) + vts.get("advancedSlicerVisual", 0)
        pgrows.append(
            [pg.get("display_name", "?"), pg.get("type", "Standard"), str(tot), str(dv), str(sl), str(tot - dv - sl)]
        )
    data_table(
        doc,
        ["Page", "Page Type", "Total Visuals", "Data Visuals", "Slicers", "Decorative"],
        pgrows,
        [2400, 1400, 800, 800, 1000, 1000],
        compact=True,
        repeat_header=True,
        title="Report structure",
    )

    # Build lookup maps for action resolution
    pages_map = {pg.get("id", ""): pg.get("display_name", pg.get("id", "")) for pg in pages}
    bookmarks_map = {bm.get("id", ""): bm.get("display_name", bm.get("id", "")) for bm in rpt.get("bookmarks", [])}
    visuals_map: dict[typing.Any, typing.Any] = {}
    for pg in pages:
        p_name = pg.get("display_name", pg.get("id", ""))
        p_id = pg.get("id", "")
        for vis in pg.get("visuals", []):
            v_id = vis.get("id", "")
            if v_id:
                visuals_map[v_id] = {
                    "visual_type": vis.get("visual_type", "?"),
                    "button_text": vis.get("button_text", ""),
                    "page_name": p_name,
                    "page_id": p_id,
                }

    # The row pipeline is intentionally staged so that page-break decisions are
    # made before the heading of each block is rendered, not after.
    _page_h = usable_page_height_in(doc)
    _y = 0.0  # cursor: estimated inches consumed on the current page

    fig_num = 3
    _prev_compact = False
    _COMPACT_VC = 4

    def _is_compact(pg: typing.Any) -> typing.Any:
        return pg.get("visual_count", len(pg.get("visuals", []))) <= _COMPACT_VC or (
            pg.get("type", "") or ""
        ).lower() in ("tooltip",)

    for pg_idx, pg in enumerate(pages):
        is_compact = _is_compact(pg)
        pg_type = pg.get("type", "Standard")
        pg_name = pg.get("display_name", "Page")
        pg_label = (
            f"Report page \u2014 {pg_name}{(' [' + pg_type + ']') if pg_type and pg_type.lower() != 'standard' else ''}"
        )
        visuals = pg.get("visuals", [])
        slicers = [v for v in visuals if v.get("visual_type", "") in ("slicer", "advancedSlicerVisual")]
        dvis = [
            v
            for v in visuals
            if v.get("visual_type", "")
            not in ("slicer", "advancedSlicerVisual", "shape", "basicShape", "image", "textbox")
        ]

        # --- PAGE BREAK DECISION (before heading is rendered) ---
        # We decide the page break before rendering the heading because compact
        # blocks such as wireframes and small tables should stay with their
        # heading whenever they fit on one page.
        # Dry-run gives exact canvas height using the real layout geometry —
        # a rough estimate would be too inaccurate for tall wireframes with
        # table rows (1.2–3.5") or multi-zone layouts.
        wf_exact_emu = insert_page_layout(doc, pg, fig_num, dry_run=True)
        block_h = estimate_wireframe_block_h(wf_exact_emu)

        if should_page_break(_page_h - _y, block_h, _page_h):
            # Block fits on a page but not in the remaining space: break first.
            page_break(doc)
            _y = 0.0
            _prev_compact = is_compact
        elif not is_compact:
            page_break(doc)
            _y = 0.0
            _prev_compact = False
        elif not _prev_compact or pg_idx == 0:
            page_break(doc)
            _y = 0.0
            _prev_compact = True
        else:
            p_div = doc.add_paragraph()
            p_div.paragraph_format.space_before = Pt(10)
            p_div.paragraph_format.space_after = Pt(0)
            _run(p_div, "\u2500" * 60, size=7, color=C.BLIGHT)
            _prev_compact = True
            _y += 0.10

        # --- RENDER REPORT-PAGE SECTION ---
        # keep_with_next chains h3 → h4 → wireframe so Word doesn't orphan
        # the headings if the diagram still doesn't fit after the break.
        p_h3 = h3(doc, pg_label)
        p_h3.paragraph_format.keep_with_next = True
        _y += _H3_IN

        # The wireframe height is dynamic, so the cursor is updated with the
        # actual rendered height returned by insert_page_layout().
        p_h4 = h4(doc, "Wireframe")
        p_h4.paragraph_format.keep_with_next = True
        _y += _H4_IN
        actual_wf_h = insert_page_layout(doc, pg, fig_num)
        _y += actual_wf_h / 914400 + _BLOCK_BUFFER_IN
        fig_num += 1

        if slicers:
            frows: list[typing.Any] = []
            for s in slicers:
                parts = [
                    f.get("property", "")
                    for f in s.get("fields", [])
                    if f.get("property", "") and f.get("property", "") != "."
                ] or ["\u2014"]
                ftype = s.get("slicer_mode") or (
                    s.get("filters", [{}])[0].get("type", "\u2014") if s.get("filters") else "\u2014"
                )
                for part in parts:
                    frows.append([part, ftype])
            # Small tables should stay with their heading when they fit on one page.
            tbl_h = estimate_table_block_h(len(frows), compact=True, has_heading=True)
            if should_page_break(_page_h - _y, tbl_h, _page_h):
                page_break(doc)
                _y = 0.0
            p_h4_f = h4(doc, "Filters")
            p_h4_f.paragraph_format.keep_with_next = True
            data_table(doc, ["Filter Field", "Filter Type"], frows, [5534, 4500], compact=True)
            _y += tbl_h

        if dvis:
            dvr: list[typing.Any] = []
            for v in dvis:
                vtype = v.get("visual_type", "?")
                if not vtype or vtype == "(not found)" or len(vtype) > 40:
                    vtype = "(custom visual)"
                fld_parts = [
                    f"{f.get('entity', '')}.{f.get('property', '')}{('(' + f.get('role', '') + ')') if f.get('role') else ''}"
                    for f in v.get("fields", [])
                    if f.get("entity", "") != "." and f.get("property", "") != "."
                ]
                flds = "; ".join(fld_parts)
                dvr.append([vtype, (flds[:70] if mode != "full" else flds) or "\u2014", v.get("id", "")[:12]])
            tbl_h = estimate_table_block_h(len(dvr), compact=True, has_heading=True)
            if should_page_break(_page_h - _y, tbl_h, _page_h):
                page_break(doc)
                _y = 0.0
            p_h4_d = h4(doc, "Data Visuals")
            p_h4_d.paragraph_format.keep_with_next = True
            data_table(doc, ["Visual Type", "Data Bindings", "Visual ID"], dvr, [1800, 6234, 2000], compact=True)
            _y += tbl_h

        vi = pg.get("visual_interactions", [])
        if vi:
            vs: dict[typing.Any, typing.Any] = {}
            for r in vi:
                vs[r.get("type", "?")] = vs.get(r.get("type", "?"), 0) + 1
            callout(
                doc,
                f"Interaction overrides: {', '.join(f'{k}: {v}' for k, v in vs.items())}."
                + (" Detail in Full mode." if mode != "full" else ""),
            )
            if mode == "full":
                tm = {v.get("id", ""): v.get("visual_type", "?") for v in visuals}
                data_table(
                    doc,
                    ["Source", "Target", "Rule"],
                    [
                        [
                            f"{tm.get(r.get('source', ''), '?')} ({r.get('source', '')[:8]})",
                            f"{tm.get(r.get('target', ''), '?')} ({r.get('target', '')[:8]})",
                            r.get("type", "?"),
                        ]
                        for r in vi
                    ],
                    [3500, 3500, 3034],
                )
        # Render visual actions / navigation links
        _render_visual_actions(doc, visuals, pages_map, bookmarks_map, mode)

        if pg.get("page_binding_type"):
            h4(doc, f"[Drillthrough]  {pg_name}  ({pg['page_binding_type']} page)")

    # Bookmarks section (report-scoped, after all pages)
    bookmarks = rpt.get("bookmarks", [])
    if bookmarks:
        page_break(doc)
        _render_bookmarks_section(doc, bookmarks, mode, pages_map, visuals_map)

    # Custom visuals section (report-scoped)
    custom_visuals = rpt.get("public_custom_visuals", [])
    if custom_visuals:
        page_break(doc)
        _render_custom_visuals_section(doc, custom_visuals)

    # Data Functions section (report-scoped)
    data_functions = _collect_data_functions(pages)
    if data_functions:
        page_break(doc)
        _render_data_functions_section(doc, data_functions, mode)

    page_break(doc)
    # 3.5 Model security (RLS) — dynamic rendering based on extracted roles
    roles = sm.get("roles", [])
    role_count = sm.get("role_count", len(roles))
    if roles:
        h2(doc, "3.5 Data security")
        bullet_item(doc, f"Row Level Security is implemented with {role_count} role(s)")

        personal_name_roles = [r.get("name") for r in roles if r.get("looks_like_personal_name") and r.get("name")]
        if personal_name_roles:
            callout(
                doc,
                f"[Warning: Roles appear to be personal names rather than functional roles: {', '.join(personal_name_roles)}. Consider using functional role names like 'Manager', 'Admin' for better maintainability.]",
                kind="warn",
            )

        rows: list[typing.Any] = []
        for idx, role in enumerate(roles, start=1):
            role_name = (
                (role.get("name") or "").strip() or Path(role.get("file", "")).stem.replace("_", " ") or f"Role {idx}"
            )
            model_perm = role.get("model_permission", "read")
            filtered_tables = role.get("tables", []) or [
                tp.get("table") for tp in role.get("table_permissions", []) if tp.get("table")
            ]
            if filtered_tables:
                sample_tables = ", ".join(str(t) for t in filtered_tables[:4])
                suffix = "..." if len(filtered_tables) > 4 else ""
                table_summary = f"{len(filtered_tables)} table(s): {sample_tables}{suffix}"
            else:
                table_summary = "All data (no table-level filters)"
            if role.get("looks_like_personal_name"):
                role_name = f"{role_name} ⚠️"
            rows.append([role_name, model_perm, table_summary])
        data_table(doc, ["Role Name", "Permission", "Tables / filter scope"], rows, [3200, 1200, 6534])
        if mode == "full":
            _render_rls_details(doc, roles)
    else:
        h2(doc, "3.5 Data security")
        body(doc, "No RLS roles defined in the semantic model. All authorized users see the full dataset.")
        if report_name.startswith("MSL"):
            callout(doc, "[Note: MSL reports typically require RLS. Verify if this is intentional.]")

    page_break(doc)
    h2(doc, "3.6 Technical observations")
    obs_list = generate_observations(sm, rpt)
    obs_list.extend(_naming_convention_observations(sm))
    limit = len(obs_list) if mode == "full" else min(12, len(obs_list))
    _OBS_CARD_COUNTER[0] = 0
    cats: dict[typing.Any, typing.Any] = {}
    for title, detail, accent, cat in obs_list[:limit]:
        cats.setdefault(cat, []).append((title, detail, accent))
    cat_colors: dict[typing.Any, typing.Any] = {
        "Risks": C.RUBINE,
        "Warnings": C.MARIGOLD,
        "Info": C.SKY,
        "Good Practices": C.EVERGREEN,
    }
    for cat_name, items in cats.items():
        p_obs_h3 = h3(doc, cat_name)
        p_obs_h3.paragraph_format.keep_with_next = True
        for title, detail, _ in items:
            obs_card(doc, title, detail, cat_colors.get(cat_name, C.SKY))
        doc.add_paragraph()

    if mode == "default":
        h2(doc, "3.7 Optional appendix on request")
        for item in [
            "Detailed table/column inventory",
            "Full DAX export",
            "Full M scripts",
            "Relationship list",
            "Visual inventory by page",
        ]:
            bullet_item(doc, item)
    page_break(doc)

    h1(doc, "4. Handover and support")
    h2(doc, "4.1 Support contacts")
    prop_table(
        doc,
        [
            ("Business Owner", "[To be confirmed]"),
            ("Technical Owner", "[To be confirmed]"),
            ("Support Queue", "[To be confirmed]"),
        ],
    )

    h2(doc, "4.2 Typical support scenarios")
    suggested(doc, "Top scenarios below generated from PBIP analysis. Confirm, modify, or add.")
    dc = _detect_connectors(exprs, tables)
    fns = {fn for fn, _, _ in dc}
    cats_set = {cat for _, _, cat in dc}
    srcs = set(sm.get("data_sources", []))
    SCENARIOS: list[typing.Any] = [
        (
            "SharePoint" in cats_set or any("sharepoint" in s.lower() for s in srcs),
            [
                [
                    "Refresh failure (SharePoint)",
                    "Source list/library schema changed",
                    "Compare SharePoint schema against Power Query column references; update affected queries",
                ],
                [
                    "Missing or incomplete data",
                    "SharePoint list view threshold exceeded (>5,000 items)",
                    "Check list item count; use API pagination or filtered queries if over threshold",
                ],
            ],
        ),
        (
            "SharePoint.Files" in fns,
            [
                [
                    "SharePoint file not found",
                    "Mapping or parameter file moved, renamed, or permissions changed",
                    "Verify file path in Power Query; check SharePoint library permissions",
                ],
            ],
        ),
        (
            any(fn in fns for fn in ("Excel.Workbook", "Excel.CurrentWorkbook"))
            or any("xlsx" in str(e.get("expression", "")).lower() for e in exprs),
            [
                [
                    "Mapping file error",
                    "Excel file moved, renamed, or column structure changed",
                    "Verify file exists at expected path with correct column names",
                ],
                [
                    "New dimension values show as unmatched",
                    "New codes/IDs not yet added to mapping file",
                    "Business user adds new rows to mapping Excel file; refresh dataset",
                ],
            ],
        ),
        (
            any(fn in fns for fn in ("Sql.Database", "Sql.Databases", "AzureSQL.Database")),
            [
                [
                    "Refresh timeout (SQL)",
                    "Source query performance degraded or database maintenance window",
                    "Check SQL query execution time; coordinate with DBA",
                ],
                [
                    "Credential expiry (SQL)",
                    "Service account password rotated without updating gateway",
                    "Re-enter SQL credentials in Power BI service",
                ],
            ],
        ),
        (
            "Snowflake.Databases" in fns,
            [
                [
                    "Snowflake refresh failure",
                    "Virtual warehouse suspended or role permissions changed",
                    "Check warehouse status; verify role grants; whitelist gateway IP",
                ],
                [
                    "Snowflake query timeout",
                    "Complex query hitting credit limits or cold cache",
                    "Review query complexity; consider result caching or materialised views",
                ],
            ],
        ),
        (
            any("Databricks" in fn for fn in fns),
            [
                [
                    "Databricks cluster not running",
                    "Auto-terminating cluster shut down before scheduled refresh",
                    "Set cluster to always-on or schedule warm-up job before refresh window",
                ],
                [
                    "Databricks schema drift",
                    "Delta table schema evolved without updating Power Query",
                    "Re-run Power Query schema detection; update column references",
                ],
            ],
        ),
        (
            "OData.Feed" in fns,
            [
                [
                    "OData feed change",
                    "Service endpoint URL changed or API version deprecated",
                    "Update OData URL; check service changelog for breaking changes",
                ]
            ],
        ),
        (
            "Web.Contents" in fns,
            [
                [
                    "REST API failure",
                    "API key expired, rate limit hit, or endpoint URL changed",
                    "Refresh API credentials; check rate-limit headers; verify endpoint URL",
                ]
            ],
        ),
        (
            "Oracle.Database" in fns,
            [
                [
                    "Oracle gateway error",
                    "On-premises data gateway driver mismatch or Oracle client not installed",
                    "Verify Oracle ODAC client version on gateway machine",
                ]
            ],
        ),
        (
            any(fn in fns for fn in ("SapHana.Database", "SapBusinessWarehouse.Cubes")),
            [
                [
                    "SAP connection failure",
                    "SAP logon tickets expired or network route blocked",
                    "Re-authenticate SAP credentials; verify RFC/BAPI network path",
                ]
            ],
        ),
        (
            "PowerBI.Datasets" in fns,
            [
                [
                    "Composite model out of sync",
                    "Source Power BI dataset refreshed at a different cadence",
                    "Align refresh schedules; check dataset endorsement status",
                ]
            ],
        ),
    ]
    scenarios = [row for cond, rows in SCENARIOS if cond for row in rows]
    scenarios += [
        [
            "Stale data after expected refresh",
            "Scheduled refresh failed silently or credentials expired",
            "Check refresh history; re-enter credentials if expired",
        ],
        [
            "Access request for new user",
            "User needs access to workspace or published app",
            "Add user to Entra ID security group or app audience",
        ],
        [
            "Visual not filtering as expected",
            "Cross-filter interaction rules or report-level filters blocking",
            "Check visual interaction settings; review page and report filters",
        ],
        [
            "Performance degradation",
            "Model size growth, complex DAX, or gateway contention",
            "Monitor model size; review DAX performance; check gateway health",
        ],
    ]
    if rpt.get("dataset_mode", "Import") == "Import":
        scenarios.append(
            [
                "Data shows as of wrong date",
                "Report shows last-refresh timestamp, user expects real-time",
                "Educate user on Import mode refresh schedule",
            ]
        )
    data_table(doc, ["Scenario", "Likely Cause", "Resolution"], scenarios, [2200, 3800, 4034], compact=True)

    h2(doc, "4.3 Manual maintenance activities")
    suggested_bullet(doc, "Update region mapping file when new company codes are added")
    suggested_bullet(doc, "Provision workspace/app access for new users via Entra ID groups")
    h2(doc, "4.4 Related evidence and documents")
    placeholder(doc, "Testing evidence")
    placeholder(doc, "Release notes / change record")
    placeholder(doc, "Access document")
    page_break(doc)

    h1(doc, "5. Revision history")
    data_table(
        doc,
        ["Version", "Date", "Author", "Change"],
        [["1.0", datetime.now().strftime("%Y-%m-%d"), "[Document Owner]", "Initial version"], ["", "", "", ""]],
        [1200, 1800, 2500, 4534],
        compact=True,
    )

    _remove_unnecessary_breaks(doc)
    return doc
