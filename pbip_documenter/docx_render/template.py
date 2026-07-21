"""Template loading and header manipulation."""

import typing

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Pt

from pbip_documenter.config import C


def _clear_template_body(doc: typing.Any) -> typing.Any:
    """Remove all body content from the template while preserving headers/footers intact."""
    body = doc.element.body
    for child in list(body):
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
        if tag not in ("sdt", "sectPr"):
            body.remove(child)
    # NOTE: Header/footer drawings are intentionally NOT stripped here.
    # The previous code removed drawings containing "Proprietary" or "Organon"
    # text, which inadvertently stripped the template logo/header image.
    return doc


def _update_template_header(
    doc: typing.Any, report_name: typing.Any, doc_id: typing.Any = "", cmdb_id: typing.Any = ""
) -> typing.Any:
    for sec in doc.sections:
        hdr = sec.header
        if not hdr.tables:
            continue
        for row in hdr.tables[0].rows:
            for cell in row.cells:
                if "Specification for:" in cell.text or "Specification for" in cell.text:
                    for p in cell.paragraphs:
                        p.clear()
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(2)
                    r = p.add_run("Power BI Solution Design and Configuration Specification for:")
                    r.font.name = "Arial"
                    r.font.size = Pt(8)
                    r.font.bold = True
                    br = OxmlElement("w:br")
                    r._r.addnext(br)
                    r2 = p.add_run(report_name)
                    r2.font.name = "Arial"
                    r2.font.size = Pt(9)
                    r2.font.bold = True
                    r2.font.color.rgb = C.rgb(C.PACIFIC)
        break
