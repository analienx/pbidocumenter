"""DrawingML shape primitives for native Word diagrams."""

import html
import typing

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from lxml import etree

from pbip_documenter.config import C

DML_NS: dict[typing.Any, typing.Any] = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
    "wpg": "http://schemas.microsoft.com/office/word/2010/wordprocessingGroup",
    "wpc": "http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas",
}


def _dns(p: typing.Any, l: typing.Any) -> typing.Any:
    return f"{{{DML_NS[p]}}}{l}"


def _dsub(parent: typing.Any, p: typing.Any, l: typing.Any, **a: typing.Any) -> typing.Any:
    e = etree.SubElement(parent, _dns(p, l))
    for k, v in a.items():
        e.set(k, str(v))
    return e


def _emu(inches: typing.Any) -> typing.Any:
    return int(inches * 914400)


# DRAWINGML ID MANAGEMENT
# Word DrawingML objects need unique IDs inside the final .docx.
# Shape IDs are used by boxes; connector IDs are used by relationship lines;
# diagram canvas docPr IDs are used by Word for each inline drawing.
# Duplicate IDs can make snapping metadata work only for the first diagram.
_AUTO_SHAPE_REGISTRY: list[typing.Any] = []
_AUTO_SHAPE_ID = 50000
_AUTO_CONNECTOR_ID = 80000
_AUTO_DIAGRAM_ID = 90000


def _next_auto_shape_id() -> typing.Any:
    global _AUTO_SHAPE_ID
    _AUTO_SHAPE_ID += 1
    return _AUTO_SHAPE_ID


def _next_auto_connector_id() -> typing.Any:
    global _AUTO_CONNECTOR_ID
    _AUTO_CONNECTOR_ID += 1
    return _AUTO_CONNECTOR_ID


def _next_auto_diagram_id() -> typing.Any:
    global _AUTO_DIAGRAM_ID
    _AUTO_DIAGRAM_ID += 1
    return _AUTO_DIAGRAM_ID


# TWEAK: connector snapping fallback
# When a caller does not pass explicit start/end shape IDs, this function chooses the
# closest registered shape and connection side. Schema diagrams pass explicit IDs, but
# lineage-style diagrams can use this fallback.
def _nearest_shape_id(x: typing.Any, y: typing.Any) -> typing.Any:
    if not _AUTO_SHAPE_REGISTRY:
        return None, None
    best = min(_AUTO_SHAPE_REGISTRY, key=lambda it: (it["x"] + it["w"] / 2 - x) ** 2 + (it["y"] + it["h"] / 2 - y) ** 2)
    cx = best["x"] + best["w"] / 2
    cy = best["y"] + best["h"] / 2
    dx = x - cx
    dy = y - cy
    idx = ("3" if dx >= 0 else "1") if abs(dx) >= abs(dy) else ("2" if dy >= 0 else "0")
    return best["id"], idx


def _escape_dml(s: typing.Any) -> typing.Any:
    return (
        html.unescape(str(s or ""))
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def _box_edge_point(
    cx: typing.Any, cy: typing.Any, hw: typing.Any, hh: typing.Any, tx: typing.Any, ty: typing.Any
) -> typing.Any:
    dx, dy = tx - cx, ty - cy
    if dx == 0 and dy == 0:
        return cx, cy
    tc: list[typing.Any] = []
    if dx:
        tc.append(hw / abs(dx))
    if dy:
        tc.append(hh / abs(dy))
    t = min(tc) if tc else 1
    return int(cx + dx * t), int(cy + dy * t)


def _dml_shape(
    x: typing.Any,
    y: typing.Any,
    w: typing.Any,
    h: typing.Any,
    fill: typing.Any,
    lines: typing.Any,
    border_color: typing.Any = None,
    border_w: typing.Any = "0",
    geom: typing.Any = "roundRect",
    adj: typing.Any = "9000",
    lIns: typing.Any = "24000",
    rIns: typing.Any = "24000",
    tIns: typing.Any = "11000",
    bIns: typing.Any = "11000",
    anchor: typing.Any = "ctr",
    autofit: typing.Any = False,
    alpha_pct: typing.Any = None,
    shape_id: typing.Any = None,
    shape_name: typing.Any = None,
) -> typing.Any:
    sp = etree.Element(_dns("wps", "wsp"))
    # Keep the original WordprocessingShape model, with optional IDs for connector references.
    if shape_id is None:
        shape_id = _next_auto_shape_id()
        shape_name = shape_name or f"Shape {shape_id}"
    _AUTO_SHAPE_REGISTRY.append({"id": int(shape_id), "x": int(x), "y": int(y), "w": int(w), "h": int(h)})
    if shape_id is not None:
        _dsub(sp, "wps", "cNvPr", id=str(shape_id), name=shape_name or f"Shape {shape_id}")
        _dsub(sp, "wps", "cNvSpPr")
    else:
        _dsub(sp, "wps", "cNvSpPr")
    spPr = _dsub(sp, "wps", "spPr")
    xfrm = _dsub(spPr, "a", "xfrm")
    _dsub(xfrm, "a", "off", x=str(x), y=str(y))
    _dsub(xfrm, "a", "ext", cx=str(w), cy=str(h))
    g = _dsub(spPr, "a", "prstGeom", prst=geom)
    avL = _dsub(g, "a", "avLst")
    if adj:
        _dsub(avL, "a", "gd", name="adj", fmla=f"val {adj}")
    if fill:
        clr = _dsub(_dsub(spPr, "a", "solidFill"), "a", "srgbClr", val=fill)
        if alpha_pct is not None:
            _dsub(clr, "a", "alpha", val=str(int(alpha_pct * 1000)))
    else:
        _dsub(spPr, "a", "noFill")
    ln = _dsub(spPr, "a", "ln", w=border_w)
    if border_color:
        _dsub(_dsub(ln, "a", "solidFill"), "a", "srgbClr", val=border_color)
    else:
        _dsub(ln, "a", "noFill")
    if lines:
        tc = _dsub(_dsub(sp, "wps", "txbx"), "w", "txbxContent")
        for line in lines:
            p = _dsub(tc, "w", "p")
            pPr = _dsub(p, "w", "pPr")
            _dsub(pPr, "w", "jc", **{_dns("w", "val"): "center"})
            _dsub(pPr, "w", "spacing", **{_dns("w", "after"): "0", _dns("w", "before"): "0"})
            for rd in line.get("runs") or [line]:
                r = _dsub(p, "w", "r")
                rPr = _dsub(r, "w", "rPr")
                _dsub(rPr, "w", "rFonts", **{_dns("w", "ascii"): "Arial", _dns("w", "hAnsi"): "Arial"})
                sz = max(rd.get("sz", 16), 8)
                _dsub(rPr, "w", "sz", **{_dns("w", "val"): str(sz)})
                _dsub(rPr, "w", "szCs", **{_dns("w", "val"): str(sz)})
                if rd.get("bold"):
                    _dsub(rPr, "w", "b")
                if rd.get("italic"):
                    _dsub(rPr, "w", "i")
                _dsub(rPr, "w", "color", **{_dns("w", "val"): rd.get("color", "FFFFFF")})
                t = _dsub(r, "w", "t")
                t.set(_dns("w", "space"), "preserve")
                t.text = html.unescape(rd.get("text", ""))
    bPr = _dsub(sp, "wps", "bodyPr", anchor=anchor, anchorCtr="0", lIns=lIns, rIns=rIns, tIns=tIns, bIns=bIns)
    if autofit:
        _dsub(bPr, "a", "spAutoFit")
    return sp


def _dml_join_label(x: typing.Any, y: typing.Any, w: typing.Any, h: typing.Any, text: typing.Any) -> typing.Any:
    sp = etree.Element(_dns("wps", "wsp"))
    _dsub(sp, "wps", "cNvSpPr", txBox="1")
    spPr = _dsub(sp, "wps", "spPr")
    xfrm = _dsub(spPr, "a", "xfrm")
    _dsub(xfrm, "a", "off", x=str(x), y=str(y))
    _dsub(xfrm, "a", "ext", cx=str(w), cy=str(h))
    g = _dsub(spPr, "a", "prstGeom", prst="rect")
    _dsub(g, "a", "avLst")
    _dsub(_dsub(_dsub(spPr, "a", "solidFill"), "a", "srgbClr", val="FFFFFF"), "a", "alpha", val="50000")
    _dsub(_dsub(_dsub(spPr, "a", "ln", w="6350"), "a", "solidFill"), "a", "srgbClr", val="AAAAAA")
    tc = _dsub(_dsub(sp, "wps", "txbx"), "w", "txbxContent")
    p = _dsub(tc, "w", "p")
    pPr = _dsub(p, "w", "pPr")
    _dsub(pPr, "w", "jc", **{_dns("w", "val"): "center"})
    _dsub(pPr, "w", "spacing", **{_dns("w", "after"): "0", _dns("w", "before"): "0"})
    r = _dsub(p, "w", "r")
    rPr = _dsub(r, "w", "rPr")
    _dsub(rPr, "w", "rFonts", **{_dns("w", "ascii"): "Arial", _dns("w", "hAnsi"): "Arial"})
    for tag, val in [("sz", "14"), ("szCs", "14"), ("color", "4A4A4A")]:
        _dsub(rPr, "w", tag, **{_dns("w", "val"): val})
    _dsub(rPr, "w", "i")
    t = _dsub(r, "w", "t")
    t.text = _escape_dml(text)
    _dsub(sp, "wps", "bodyPr", anchor="ctr", anchorCtr="0", lIns="18000", rIns="18000", tIns="0", bIns="0")
    return sp


def _dml_header_label(
    x: typing.Any, y: typing.Any, w: typing.Any, h: typing.Any, text: typing.Any, color: typing.Any
) -> typing.Any:
    sp = etree.Element(_dns("wps", "wsp"))
    _dsub(sp, "wps", "cNvSpPr", txBox="1")
    spPr = _dsub(sp, "wps", "spPr")
    xfrm = _dsub(spPr, "a", "xfrm")
    _dsub(xfrm, "a", "off", x=str(x), y=str(y))
    _dsub(xfrm, "a", "ext", cx=str(w), cy=str(h))
    g = _dsub(spPr, "a", "prstGeom", prst="rect")
    _dsub(g, "a", "avLst")
    _dsub(spPr, "a", "noFill")
    _dsub(_dsub(spPr, "a", "ln", w="0"), "a", "noFill")
    tc = _dsub(_dsub(sp, "wps", "txbx"), "w", "txbxContent")
    p = _dsub(tc, "w", "p")
    pPr = _dsub(p, "w", "pPr")
    _dsub(pPr, "w", "jc", **{_dns("w", "val"): "center"})
    _dsub(pPr, "w", "spacing", **{_dns("w", "after"): "0", _dns("w", "before"): "0"})
    r = _dsub(p, "w", "r")
    rPr = _dsub(r, "w", "rPr")
    _dsub(rPr, "w", "rFonts", **{_dns("w", "ascii"): "Arial", _dns("w", "hAnsi"): "Arial"})
    for tag, val in [("sz", "20"), ("szCs", "20"), ("color", color)]:
        _dsub(rPr, "w", tag, **{_dns("w", "val"): val})
    _dsub(rPr, "w", "b")
    t = _dsub(r, "w", "t")
    t.text = _escape_dml(text)
    _dsub(sp, "wps", "bodyPr", anchor="ctr", anchorCtr="0", lIns="0", rIns="0", tIns="0", bIns="0")
    return sp


def _dml_connector(
    x1: typing.Any,
    y1: typing.Any,
    x2: typing.Any,
    y2: typing.Any,
    color: typing.Any = "888888",
    dashed: typing.Any = False,
    parallel_index: typing.Any = 0,
    start_arrow: typing.Any = False,
    end_arrow: typing.Any = True,
    route: typing.Any = "straight",
    connector_id: typing.Any = None,
    start_shape_id: typing.Any = None,
    end_shape_id: typing.Any = None,
    start_idx: typing.Any = "3",
    end_idx: typing.Any = "1",
) -> typing.Any:
    dx, dy = x2 - x1, y2 - y1
    if parallel_index:
        mag = max((dx * dx + dy * dy) ** 0.5, 1)
        px, py = -dy / mag, dx / mag
        offset = 9000 * parallel_index
        x1, y1 = int(x1 + px * offset), int(y1 + py * offset)
        x2, y2 = int(x2 + px * offset), int(y2 + py * offset)
    if connector_id is None:
        connector_id = _next_auto_connector_id()
    if start_shape_id is None:
        start_shape_id, auto_idx = _nearest_shape_id(x1, y1)
        if auto_idx is not None:
            start_idx = auto_idx
    if end_shape_id is None:
        end_shape_id, auto_idx = _nearest_shape_id(x2, y2)
        if auto_idx is not None:
            end_idx = auto_idx
    sp = etree.Element(_dns("wps", "wsp"))
    if connector_id is not None:
        _dsub(sp, "wps", "cNvPr", id=str(connector_id), name=f"Connector {connector_id}")
    cnv = _dsub(sp, "wps", "cNvCnPr")
    # Snappy-inspired hints. These are harmless if Word ignores them, and match the validated 3 -> 1 pattern where applicable.
    if start_shape_id is not None:
        _dsub(cnv, "a", "stCxn", id=str(start_shape_id), idx=str(start_idx))
    if end_shape_id is not None:
        _dsub(cnv, "a", "endCxn", id=str(end_shape_id), idx=str(end_idx))
    spPr = _dsub(sp, "wps", "spPr")
    xfrm = _dsub(spPr, "a", "xfrm")
    if x2 < x1:
        xfrm.set("flipH", "1")
    if y2 < y1:
        xfrm.set("flipV", "1")
    _dsub(xfrm, "a", "off", x=str(min(x1, x2)), y=str(min(y1, y2)))
    _dsub(xfrm, "a", "ext", cx=str(abs(x2 - x1) or 1000), cy=str(abs(y2 - y1) or 1000))
    g = _dsub(spPr, "a", "prstGeom", prst="straightConnector1" if route == "straight" else "curvedConnector3")
    _dsub(g, "a", "avLst")
    ln = _dsub(spPr, "a", "ln", w="12700")
    _dsub(_dsub(ln, "a", "solidFill"), "a", "srgbClr", val=color)
    if dashed:
        _dsub(ln, "a", "prstDash", val="sysDash")
    if end_arrow:
        _dsub(ln, "a", "headEnd", type="triangle", w="med", len="med")
    if start_arrow:
        _dsub(ln, "a", "tailEnd", type="triangle", w="med", len="med")
    _dsub(sp, "wps", "bodyPr")
    return sp


# TWEAK: diagram canvas insertion
# `w` and `h` are DrawingML canvas dimensions in EMUs.
# Schema canvas height is intentionally kept below the physical page budget in schema.py.
# Increasing the canvas too much can create blank/orphan pages even when XML compiles.
def _insert_diagram(
    doc: typing.Any, shapes: typing.Any, w: typing.Any, h: typing.Any, caption: typing.Any = None
) -> typing.Any:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    # Do not force keep_with_next here. Large inline diagrams plus captions can trigger blank pages.
    r = etree.SubElement(p._p, _dns("w", "r"))
    dr = etree.SubElement(r, _dns("w", "drawing"))
    inl = etree.SubElement(dr, _dns("wp", "inline"))
    for a in ("distT", "distB", "distL", "distR"):
        inl.set(a, "0")
    etree.SubElement(inl, _dns("wp", "extent"), cx=str(w), cy=str(h))
    ee = etree.SubElement(inl, _dns("wp", "effectExtent"))
    for a in ("l", "t", "r", "b"):
        ee.set(a, "0")
    etree.SubElement(inl, _dns("wp", "docPr"), id=str(_next_auto_diagram_id()), name="Diagram")
    gd = etree.SubElement(etree.SubElement(inl, _dns("a", "graphic")), _dns("a", "graphicData"))
    gd.set("uri", DML_NS["wpg"])
    wgp = etree.SubElement(gd, _dns("wpg", "wgp"))
    etree.SubElement(wgp, _dns("wpg", "cNvGrpSpPr"))
    gsp = etree.SubElement(wgp, _dns("wpg", "grpSpPr"))
    xf = etree.SubElement(gsp, _dns("a", "xfrm"))
    for tag, vals in [
        ("off", {"x": "0", "y": "0"}),
        ("ext", {"cx": str(w), "cy": str(h)}),
        ("chOff", {"x": "0", "y": "0"}),
        ("chExt", {"cx": str(w), "cy": str(h)}),
    ]:
        e = etree.SubElement(xf, _dns("a", tag))
        for k, v in vals.items():
            e.set(k, v)
    for s in shapes:
        wgp.append(s)
    _AUTO_SHAPE_REGISTRY.clear()
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(1)
        cp.paragraph_format.space_after = Pt(2)
        cr = cp.add_run(caption)
        cr.font.name = "Arial"
        cr.font.size = Pt(8)
        cr.font.italic = True
        cr.font.color.rgb = C.rgb(C.DGRAY)


def _diagram_legend(doc: typing.Any, items: typing.Any) -> typing.Any:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for label, color in items:
        r = p.add_run(label + "  ")
        r.font.name = "Arial"
        r.font.size = Pt(9)
        r.font.color.rgb = C.rgb(color)


def _min_text_height(
    lines: typing.Any, tIns: typing.Any = 9000, bIns: typing.Any = 9000, gap: typing.Any = 1.25
) -> typing.Any:
    total = tIns + bIns
    for line in lines or []:
        total += int((line.get("sz", 16) / 2.0) * 12700 * gap)
    return total
