"""Native Word internal navigation, available before updating TOC fields."""

from __future__ import annotations

from typing import Any

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from pbip_documenter.config import C

# The main document uses exactly these Heading 1 labels.
NAVIGATION = (
    ("Document control", "Document owner, classification and version", "toc_control"),
    ("1. Purpose and scope", "Purpose, users, scope and assumptions", "toc_scope"),
    ("2. Solution and operations summary", "Dependencies, connectivity and access", "toc_operations"),
    ("3. Technical specifications", "Model, DAX, lineage, report pages and diagrams", "toc_technical"),
    ("4. Handover and support", "Support procedures and ownership", "toc_support"),
    ("5. Revision history", "Document version and change record", "toc_revision"),
)


def add_navigation(doc: Any) -> None:
    """Populate a linked contents list that works without refreshing Word fields."""
    for title, description, anchor in NAVIGATION:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.keep_together = True

        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("w:anchor"), anchor)
        hyperlink.set(qn("w:history"), "1")

        run = OxmlElement("w:r")
        props = OxmlElement("w:rPr")
        bold = OxmlElement("w:b")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), C.PACIFIC)
        props.append(bold)
        props.append(color)
        run.append(props)
        text = OxmlElement("w:t")
        text.text = title
        run.append(text)
        hyperlink.append(run)
        p._p.append(hyperlink)

        detail = p.add_run("    " + description)
        detail.font.size = Pt(9)
        detail.font.color.rgb = C.rgb(C.DGRAY)


def bookmark_navigation_targets(doc: Any) -> None:
    """Anchor the actual Heading 1 paragraphs to the linked contents entries."""
    targets = {title: anchor for title, _, anchor in NAVIGATION}
    found: set[str] = set()
    for p in doc.paragraphs:
        anchor = targets.get(p.text)
        if anchor is None or p.style.name != "Heading 1":
            continue
        if anchor in found:
            raise ValueError(f"Duplicate contents heading: {p.text}")
        found.add(anchor)
        marker_id = str(15000 + len(found))

        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), marker_id)
        start.set(qn("w:name"), anchor)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), marker_id)

        p._p.insert(1 if p._p.pPr is not None else 0, start)
        p._p.append(end)

    missing = set(targets.values()) - found
    if missing:
        raise ValueError(f"Contents targets missing from Word document: {sorted(missing)}")
